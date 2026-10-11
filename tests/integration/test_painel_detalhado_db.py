"""Daily and hierarchy acceptance through signed HTTP, PostgreSQL and forced RLS."""

import io
import time
from datetime import UTC, date, datetime, timedelta
from itertools import pairwise
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy import event, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from bancaemdia import models
from bancaemdia.config import get_settings
from bancaemdia.domain import painel_cursor
from bancaemdia.domain.materializar import casa_canonica
from tests.integration.test_site_filters_db import api as api
from tests.integration.test_site_filters_db import dataset as dataset

pytestmark = pytest.mark.xdist_group("postgres")
PREFIX = "/api/v1/painel/filtrado"


@pytest.fixture
async def detailed(api: Any, dataset: Any, engine_app: Any, como: Any) -> Any:
    client, headers = api
    auth = headers(dataset.user)
    assert (
        await client.patch(
            f"/api/v1/grupos/{dataset.grupos[1]}",
            json={"nome": "Grupo 1", "arquivado": False},
            headers=auth,
        )
    ).status_code == 200
    # The same tipster really has facts in two groups; no current membership inference.
    links = [[0], [0], [0], [0], [1], [1], [], [0, 1], [], []]
    for key, indices in zip(dataset.keys, links, strict=False):
        result = await client.put(
            f"/api/v1/apostas/{key}/grupos",
            headers=auth,
            json={"grupo_ids": [str(dataset.grupos[i]) for i in indices]},
        )
        assert result.status_code == 200, result.text
    assert (
        await client.patch(
            f"/api/v1/grupos/{dataset.grupos[1]}",
            json={"nome": "Grupo 1", "arquivado": True},
            headers=auth,
        )
    ).status_code == 200
    async with como(engine_app, dataset.user) as session:
        for i, key in enumerate(dataset.keys[:10]):
            row = await session.scalar(select(models.Aposta).where(models.Aposta.chave == key))
            assert row is not None
            row.data_aposta = (
                datetime(2025, 12, 31, 2, 59, 59, tzinfo=UTC)
                if i == 0
                else datetime(2025, 12, 31, 3, tzinfo=UTC) + timedelta(days=i - 1)
            )
        await session.commit()
    return dataset


async def read(api: Any, dataset: Any, suffix: Any, params: Any = None) -> Any:
    client, headers = api
    response = await client.get(PREFIX + suffix, headers=headers(dataset.user), params=params)
    assert response.status_code == 200, response.text
    assert response.headers["cache-control"] == "private, no-store"
    assert "Cookie" in response.headers["vary"]
    return response.json()


async def test_hierarchy_historical_identity_and_full_group_totals(api: Any, detailed: Any) -> None:
    body = await read(api, detailed, "/grupos")
    assert body["total_grupos"] == 4
    assert body["total_filhos"] == 7
    total = body["total_periodo"]
    assert total["total_apostas"] == 10
    assert total["lucro_centavos"] is None
    assert total["lucro_conhecido_centavos"] == "325"
    assert sum(g["metricas"]["total_apostas"] for g in body["grupos"]) == 10
    assert sum(int(g["metricas"]["lucro_conhecido_centavos"]) for g in body["grupos"]) == 325
    g0, g1 = body["grupos"][:2]
    assert g0["id"] == str(detailed.grupos[0])
    assert g0["total_filhos"] == 2
    assert g0["metricas"]["lucro_conhecido_centavos"] == "0"
    assert g0["metricas"]["lucro_centavos"] is None
    assert g1["ativo"] is False  # Archived does not erase historical facts.
    assert g1["metricas"]["lucro_centavos"] == "350"
    assert {t["id"] for t in g0["tipsters"]} == {t["id"] for t in g1["tipsters"]}
    assert body["grupos"][2]["atribuicao"] == "sem_grupo"
    assert body["grupos"][3]["atribuicao"] == "ambigua"
    assert body["grupos"][3]["metricas"]["lucro_centavos"] == "25"
    assert body["politica_atribuicao"] == "vinculo_explicito_unico_da_aposta"
    client, headers = api
    renamed = await client.patch(
        f"/api/v1/grupos/{detailed.grupos[0]}",
        json={"nome": "Mesmo rótulo do tipster"},
        headers=headers(detailed.user),
    )
    assert renamed.status_code == 200
    after = await read(api, detailed, "/grupos")
    assert after["grupos"][0]["id"] == g0["id"]
    assert after["grupos"][0]["nome"] == "Mesmo rótulo do tipster"
    assert after["grupos"][0]["metricas"] == g0["metricas"]


async def test_pairs_pagination_visits_every_child_without_changing_totals(
    api: Any, detailed: Any
) -> None:
    params = {"page_size": 1}
    pages = []
    seen = set()
    while True:
        body = await read(api, detailed, "/grupos", params)
        pages.append(body)
        group = body["grupos"][0]
        child = group["tipsters"][0]
        identity = (group["atribuicao"], group["id"], child["id"], child["atribuicao"])
        assert identity not in seen
        seen.add(identity)
        assert body["total_periodo"]["lucro_conhecido_centavos"] == "325"
        if not body["proximo_cursor"]:
            break
        params["cursor"] = body["proximo_cursor"]
    assert len(pages) == 7
    assert pages[0]["grupos"][0]["metricas"] == pages[1]["grupos"][0]["metricas"]
    previous = await read(
        api, detailed, "/grupos", {"page_size": 1, "cursor": pages[1]["anterior_cursor"]}
    )
    assert previous["grupos"] == pages[0]["grupos"]


async def test_daily_states_exact_profit_and_year_timezone_boundary(
    api: Any, detailed: Any
) -> None:
    body = await read(
        api,
        detailed,
        "/diario",
        {
            "desde": "2025-12-30T00:00:00-03:00",
            "ate": "2026-01-11T00:00:00-03:00",
        },
    )
    assert body["granularidade"] == "dia"
    assert body["criterio_temporal"] == "data_aposta"
    assert body["fuso_horario"] == "America/Sao_Paulo"
    assert body["intervalo_total"]["dias"] == 12
    values = {row["dia"]: row["metricas"] for row in body["dias"]}
    assert values["2025-12-30"]["lucro_centavos"] == "100"
    assert values["2025-12-31"]["lucro_centavos"] == "-100"
    assert values["2026-01-01"]["estado_dado"] == "pendente"
    assert values["2026-01-01"]["lucro_centavos"] is None
    assert values["2026-01-02"]["estado_dado"] == "confirmado"
    assert values["2026-01-02"]["lucro_centavos"] == "0"
    assert values["2026-01-03"]["lucro_centavos"] == "300"  # Freebet costs zero.
    assert values["2026-01-03"]["base_roi_centavos"] == "100"  # Face retained.
    assert values["2026-01-04"]["lucro_centavos"] == "50"
    assert values["2026-01-05"]["lucro_centavos"] == "-50"
    assert values["2026-01-06"]["lucro_centavos"] == "25"
    assert values["2026-01-07"]["estado_dado"] == "desconhecido"
    assert values["2026-01-08"]["estado_dado"] == "desconhecido"
    assert values["2026-01-08"]["lucro_centavos"] is None
    assert values["2026-01-09"]["estado_dado"] == "sem_atividade"
    assert values["2026-01-09"]["lucro_centavos"] == "0"
    assert body["subtotal_intervalo"] == body["total_periodo"]


@pytest.mark.parametrize("period", ["7d", "30d", "90d", "1y", "all"])
async def test_every_period_is_daily_and_fully_navigable(
    api: Any, detailed: Any, period: Any
) -> None:
    params = {"periodo": period, "dias_por_pagina": 31}
    seen: list[date] = []
    expected_total = None
    count = known = 0
    while True:
        body = await read(api, detailed, "/diario", params)
        assert body["granularidade"] == "dia"
        if expected_total is None:
            expected_total = body["total_periodo"]
        assert body["total_periodo"] == expected_total
        count += body["subtotal_intervalo"]["total_apostas"]
        known += int(body["subtotal_intervalo"]["lucro_conhecido_centavos"])
        seen.extend(date.fromisoformat(row["dia"]) for row in body["dias"])
        if not body["proximo_cursor"]:
            break
        assert len(body["dias"]) == 31
        params["cursor"] = body["proximo_cursor"]
    assert len(seen) == body["intervalo_total"]["dias"]
    assert all(b - a == timedelta(days=1) for a, b in pairwise(seen))
    assert count == expected_total["total_apostas"]
    assert known == int(expected_total["lucro_conhecido_centavos"])
    if period == "all":
        assert seen[0] == date(2025, 12, 30)
        assert len(seen) > 270
    elif period in {"7d", "30d", "90d"}:
        assert len(seen) == int(period[:-1])
    else:
        assert len(seen) in {365, 366}


@pytest.mark.parametrize(
    "params",
    [
        {"grupo_id": "g0"},
        {"conta_casa_id": "c0", "titular_id": "t0", "banca_id": "b0"},
        {"origem": "planilha"},
        {"visibilidade": "apagadas"},
        {"visibilidade": "todas"},
    ],
)
async def test_filters_and_complete_excel_reconcile(api: Any, detailed: Any, params: Any) -> None:
    ids = {
        "g0": detailed.grupos[0],
        "c0": detailed.contas[0],
        "t0": detailed.titulares[0],
        "b0": detailed.bancas[0],
    }
    params = {k: str(ids.get(v, v)) for k, v in params.items()}
    daily = await read(api, detailed, "/diario", {**params, "dias_por_pagina": 366})
    hierarchy = await read(api, detailed, "/grupos", params)
    assert daily["total_periodo"] == hierarchy["total_periodo"]
    client, headers = api
    listing = await client.get("/api/v1/apostas", params=params, headers=headers(detailed.user))
    summary = await client.get(PREFIX, params=params, headers=headers(detailed.user))
    total = daily["total_periodo"]
    assert listing.json()["pagination"]["total"] == total["total_apostas"]
    assert summary.json()["resumo"]["total_apostas"] == total["total_apostas"]
    response = await client.get(
        PREFIX + "/detalhado/export", params=params, headers=headers(detailed.user)
    )
    assert response.status_code == 200, response.text
    assert response.headers["cache-control"] == "private, no-store"
    book = load_workbook(io.BytesIO(response.content), read_only=True)
    try:
        for name in ("Grupos", "Tipsters por grupo", "Diario"):
            rows = list(book[name].values)
            count_idx = rows[0].index("total_apostas")
            known_idx = rows[0].index("lucro_conhecido_centavos")
            assert sum(r[count_idx] for r in rows[1:]) == total["total_apostas"]
            assert sum(int(r[known_idx]) for r in rows[1:]) == int(
                total["lucro_conhecido_centavos"]
            )
        assert "grupo_ativo" in next(book["Grupos"].values)
        assert "tipster_atribuicao" in next(book["Tipsters por grupo"].values)
        assert "escopo" in next(book["Contexto"].values)
        assert list(book["Resumo"].values)[1][0] == total["total_apostas"]
    finally:
        book.close()


async def test_missing_dates_no_tipster_empty_history_and_previous_window(
    api: Any, detailed: Any
) -> None:
    body = await read(api, detailed, "/diario", {"visibilidade": "apagadas"})
    assert body["intervalo_total"]["dias"] == 0
    assert body["sem_data"]["total_apostas"] == 2
    assert body["sem_data"]["lucro_centavos"] is None
    assert body["total_periodo"]["lucro_conhecido_centavos"] == "-100"
    hierarchy = await read(api, detailed, "/grupos", {"visibilidade": "apagadas"})
    assert hierarchy["grupos"][0]["tipsters"][0]["atribuicao"] == "sem_tipster"
    empty = await read(api, detailed, "/diario", {"origem": "inexistente"})
    assert empty["dias"] == []
    assert empty["total_periodo"]["estado_dado"] == "sem_atividade"
    first = await read(api, detailed, "/diario", {"periodo": "all", "dias_por_pagina": 2})
    second = await read(
        api,
        detailed,
        "/diario",
        {"periodo": "all", "dias_por_pagina": 2, "cursor": first["proximo_cursor"]},
    )
    previous = await read(
        api,
        detailed,
        "/diario",
        {"periodo": "all", "dias_por_pagina": 2, "cursor": second["anterior_cursor"]},
    )
    assert previous["dias"] == first["dias"]


async def test_cursor_privacy_tampering_cross_owner_filter_and_projection(
    api: Any, detailed: Any
) -> None:
    daily = await read(api, detailed, "/diario", {"dias_por_pagina": 1})
    cursor = daily["proximo_cursor"]
    assert cursor and detailed.keys[0] not in cursor
    client, headers = api
    for uid, suffix, params in [
        (detailed.other, "/diario", {"cursor": cursor, "dias_por_pagina": 1}),
        (detailed.user, "/grupos", {"cursor": cursor}),
        (detailed.user, "/diario", {"cursor": cursor, "dias_por_pagina": 2}),
        (detailed.user, "/diario", {"cursor": cursor, "dias_por_pagina": 1, "origem": "manual"}),
        (detailed.user, "/diario", {"cursor": "?" + cursor[1:]}),
        (detailed.user, "/diario", {"cursor": cursor, "inicio_dia": "2026-01-01"}),
        (detailed.user, "/diario", {"dias_por_pagina": 367}),
        (detailed.user, "/diario", {"inicio_dia": "2020-01-01"}),
        (detailed.user, "/grupos", {"page_size": 101}),
        (detailed.user, "/detalhado/export", {"desde": "1000-01-01T00:00:00Z"}),
        (detailed.user, "/diario", {"ate": "9999-12-31T23:59:59-03:00"}),
    ]:
        response = await client.get(PREFIX + suffix, params=params, headers=headers(uid))
        assert response.status_code == 422, response.text
        assert response.headers["cache-control"] == "private, no-store"
    for suffix in ("/diario", "/grupos", "/detalhado/export"):
        assert (await client.get(PREFIX + suffix)).status_code == 401


@pytest.mark.parametrize("suffix", ["/diario", "/grupos", "/detalhado/export"])
@pytest.mark.parametrize("key", ["grupo_id", "banca_id", "titular_id", "conta_casa_id"])
async def test_foreign_selectors_are_indistinguishable_from_missing(
    api: Any, detailed: Any, suffix: Any, key: Any
) -> None:
    client, headers = api
    foreign = await client.get(
        PREFIX + suffix,
        params={key: str(detailed.other_ids[key])},
        headers=headers(detailed.user),
    )
    absent = await client.get(
        PREFIX + suffix, params={key: str(2**63 - 1)}, headers=headers(detailed.user)
    )
    assert foreign.status_code == absent.status_code == 200
    if suffix != "/detalhado/export":
        assert foreign.json()["total_periodo"] == absent.json()["total_periodo"]
        assert foreign.json()["total_periodo"]["total_apostas"] == 0


async def test_read_only_new_reads_and_exports_remain_available(
    api: Any, detailed: Any, engine_admin: Any
) -> None:
    async with engine_admin.begin() as conn:
        previous = await conn.scalar(text("SELECT activated_at FROM billing_rollout WHERE id=1"))
        await conn.execute(text("UPDATE billing_rollout SET activated_at=now() WHERE id=1"))
    try:
        client, headers = api
        for suffix in ("/diario", "/grupos", "/detalhado/export"):
            assert (
                await client.get(PREFIX + suffix, headers=headers(detailed.user))
            ).status_code == 200
        assert (
            await client.post(
                "/api/v1/grupos", json={"nome": "Bloqueado"}, headers=headers(detailed.user)
            )
        ).status_code == 402
        assert (await client.get(PREFIX + "/diario")).status_code == 401
    finally:
        async with engine_admin.begin() as conn:
            await conn.execute(
                text("UPDATE billing_rollout SET activated_at=:v WHERE id=1"), {"v": previous}
            )


@pytest.mark.parametrize("suffix", ["/diario", "/grupos", "/detalhado/export"])
async def test_real_statement_timeout_returns_private_503(
    api: Any, detailed: Any, monkeypatch: pytest.MonkeyPatch, suffix: Any
) -> None:
    original = AsyncSession.execute

    async def timed_out(session: Any, statement: Any, *args: Any, **kwargs: Any) -> Any:
        if "apostas_selecionadas" in str(statement):
            await original(session, text("SET LOCAL statement_timeout = '1ms'"))
            return await original(session, text("SELECT pg_sleep(0.05)"))
        return await original(session, statement, *args, **kwargs)

    monkeypatch.setattr(AsyncSession, "execute", timed_out)
    client, headers = api
    response = await client.get(PREFIX + suffix, headers=headers(detailed.user))
    assert response.status_code == 503
    assert response.json() == {"detail": "projeção temporariamente indisponível"}
    assert response.headers["cache-control"] == "private, no-store"


async def test_rate_limits_and_expired_cursor_are_private(
    api: Any, detailed: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    client, headers = api
    daily = await read(api, detailed, "/diario", {"dias_por_pagina": 1})
    token = daily["proximo_cursor"]

    class Tomorrow(datetime):
        @classmethod
        def now(cls, tz: Any = None) -> Any:
            return datetime.now(tz) + timedelta(days=2)

    monkeypatch.setattr(painel_cursor, "datetime", Tomorrow)
    # No timestamp filters here, avoiding any synthetic context reserialization.
    response = await client.get(
        PREFIX + "/diario",
        params={"cursor": token, "dias_por_pagina": 1},
        headers=headers(detailed.user),
    )
    assert response.status_code == 422
    monkeypatch.setattr(painel_cursor, "datetime", datetime)
    monkeypatch.setattr(get_settings(), "PANEL_EXPORT_RATE_LIMIT", "1/minute")
    assert (
        await client.get(PREFIX + "/detalhado/export", headers=headers(detailed.user))
    ).status_code == 200
    limited = await client.get(PREFIX + "/detalhado/export", headers=headers(detailed.user))
    assert limited.status_code == 429
    assert limited.headers["cache-control"] == "private, no-store"
    assert "retry-after" in limited.headers


async def test_bigint_identity_and_money_remain_exact_in_json_and_excel(
    api: Any, detailed: Any, engine_app: Any, como: Any
) -> None:
    large = 2**53 + 1 + detailed.user
    amount = 2**53 + 1
    async with como(engine_app, detailed.user) as session:
        tipster = models.Tipster(id=large, nome=f"BIGINT {large}")
        group = models.GrupoAposta(id=large, usuario_id=detailed.user, nome=f"BIGINT {large}")
        bet = models.Aposta(
            usuario_id=detailed.user,
            chave=f"exact:{large}",
            origem="manual",
            estado="RED",
            stake_unidades=1,
            stake_centavos=amount,
            valor_aposta_centavos=amount,
            retorno_centavos=0,
            selecionada=True,
            tipster_id=large,
            data_aposta=datetime(2026, 1, 1, tzinfo=UTC),
        )
        session.add_all([tipster, group])
        await session.flush()
        session.add(bet)
        await session.flush()
        session.add(models.ApostaGrupo(usuario_id=detailed.user, aposta_id=bet.id, grupo_id=large))
        await session.commit()
    body = await read(api, detailed, "/grupos", {"grupo_id": str(large)})
    assert body["grupos"][0]["id"] == str(large)
    assert body["grupos"][0]["tipsters"][0]["id"] == str(large)
    assert body["total_periodo"]["lucro_centavos"] == str(-amount)
    client, headers = api
    response = await client.get(
        PREFIX + "/detalhado/export",
        params={"grupo_id": str(large)},
        headers=headers(detailed.user),
    )
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content), read_only=True)
    try:
        rows = list(book["Grupos"].values)
        assert rows[1][0] == str(large)
        assert rows[1][rows[0].index("lucro_centavos")] == str(-amount)
    finally:
        book.close()


async def test_long_history_query_count_is_bounded_and_groups_not_n_plus_one(
    api: Any, detailed: Any, engine_app: Any, como: Any, record_property: Any
) -> None:
    calls = []

    def count_queries(
        conn: Any, cursor: Any, statement: Any, parameters: Any, context: Any, executemany: Any
    ) -> Any:
        calls.append(statement)

    event.listen(engine_app.sync_engine, "before_cursor_execute", count_queries)
    try:
        began = time.perf_counter()
        await read(api, detailed, "/grupos")
        small = len(calls)
        calls.clear()
        async with como(engine_app, detailed.user) as session:
            # 150 additional independently assigned group facts; actual aggregate, not mocked SQL.
            for i in range(150):
                group = models.GrupoAposta(usuario_id=detailed.user, nome=f"Carga {i}")
                bet = models.Aposta(
                    usuario_id=detailed.user,
                    chave=f"load:{detailed.keys[0]}:{i}",
                    origem="manual",
                    estado="RED",
                    stake_unidades=1,
                    stake_centavos=100,
                    valor_aposta_centavos=100,
                    retorno_centavos=0,
                    selecionada=True,
                    data_aposta=datetime(2025, 1, 1, tzinfo=UTC),
                )
                session.add_all([group, bet])
                await session.flush()
                session.add(
                    models.ApostaGrupo(
                        usuario_id=detailed.user, aposta_id=bet.id, grupo_id=group.id
                    )
                )
            await session.commit()
        calls.clear()
        body = await read(api, detailed, "/grupos", {"page_size": 100})
        large = len(calls)
        assert large == small <= 14
        assert body["total_grupos"] == 154
        assert body["total_periodo"]["lucro_conhecido_centavos"] == "-14675"
        calls.clear()
        daily = await read(api, detailed, "/diario", {"dias_por_pagina": 366})
        assert len(calls) <= 12
        assert daily["intervalo_total"]["dias"] > 640
        record_property("group_queries_small", small)
        record_property("group_queries_large", large)
        record_property("daily_queries", len(calls))
        record_property("exercise_seconds", time.perf_counter() - began)
    finally:
        event.remove(engine_app.sync_engine, "before_cursor_execute", count_queries)


@pytest.mark.parametrize("telegram_state", ["GREEN", "PENDENTE"])
async def test_canonical_sources_do_not_duplicate_group_or_daily_finances(
    api: Any, detailed: Any, engine_app: Any, como: Any, telegram_state: str
) -> None:
    async with como(engine_app, detailed.user) as session:
        group = models.GrupoAposta(usuario_id=detailed.user, nome="Fontes canônicas")
        session.add(group)
        bets = []
        for origin in ("casa", "telegram"):
            bet = models.Aposta(
                usuario_id=detailed.user,
                chave=f"canonical:{origin}:{detailed.keys[0]}",
                origem=origin,
                estado="GREEN" if origin == "casa" else telegram_state,
                revisao_grave=origin == "telegram" and telegram_state == "PENDENTE",
                selecionada=True,
                chat_id=123 if origin == "telegram" else None,
                message_id=456 if origin == "telegram" else None,
                conta_casa_id=detailed.contas[0],
                stake_unidades=1,
                stake_centavos=100,
                valor_aposta_centavos=100,
                retorno_centavos=200,
                data_aposta=datetime(2026, 1, 1, tzinfo=UTC),
                tipster_id=detailed.tipsters[0],
            )
            session.add(bet)
            bets.append(bet)
        await session.flush()
        session.add(
            models.ApostaConsolidacao(
                usuario_id=detailed.user,
                casa_aposta_id=bets[0].id,
                telegram_aposta_id=bets[1].id,
                conta_casa_id=detailed.contas[0],
                estado="active",
                decisao="reviewed",
                versao="panel-detail-acceptance",
                evidencia={},
                contexto={},
                ator="test",
            )
        )
        session.add_all([
            models.ApostaGrupo(usuario_id=detailed.user, aposta_id=bet.id, grupo_id=group.id)
            for bet in bets
        ])
        gid = group.id
        await session.commit()
    for suffix in ("/diario", "/grupos"):
        body = await read(api, detailed, suffix, {"grupo_id": str(gid), "visibilidade": "todas"})
        assert body["total_periodo"]["total_apostas"] == 2
        assert body["total_periodo"]["fontes_excluidas_financeiro"] == 1
        assert body["total_periodo"]["lucro_centavos"] == "100"
        assert body["total_periodo"]["giro_centavos"] == "100"
        assert body["total_periodo"]["estado_dado"] == "confirmado"
        assert body["total_periodo"]["pendentes_financeiras"] == 0
        assert body["total_periodo"]["revisoes_financeiras"] == 0
        if suffix == "/grupos":
            assert body["grupos"][0]["tipsters"][0]["metricas"]["lucro_centavos"] == "100"
    context_only = await read(
        api,
        detailed,
        "/grupos",
        {"grupo_id": str(gid), "visibilidade": "todas", "origem": "telegram"},
    )
    assert context_only["total_periodo"]["estado_dado"] == "fora_do_financeiro"
    assert context_only["total_periodo"]["lucro_centavos"] is None
    assert context_only["total_periodo"]["lucro_conhecido_centavos"] == "0"


async def test_actual_materialization_delete_restore_keeps_same_group_fact(
    api: Any, detailed: Any, engine_app: Any, como: Any
) -> None:
    client, headers = api
    auth = headers(detailed.user)
    async with como(engine_app, detailed.user) as session:
        house = casa_canonica("betano")
        assert house is not None
        catalog = await session.scalar(select(models.Casa).where(models.Casa.nome == house))
        if catalog is None:
            catalog = models.Casa(nome=house)
            session.add(catalog)
            await session.flush()
        account = models.ContaCasa(
            usuario_id=detailed.user,
            casa_id=catalog.id,
            titular_id=detailed.titulares[0],
            banca_id=detailed.bancas[0],
            apelido="Conta do aceite materializado",
        )
        session.add(account)
        await session.commit()
        account_id = account.id
    created = await client.post(
        "/api/v1/apostas",
        headers=auth,
        json={
            "casa": house,
            "odd": 2,
            "stake_unidades": 1,
            "data_aposta": "2026-01-01T12:00:00-03:00",
            "data_jogo": "2026-01-02T12:00:00-03:00",
            "conta_casa_ref": account_id,
        },
    )
    assert created.status_code == 201, created.text
    key = created.json()["aposta"]["chave"]
    assert (
        await client.put(
            f"/api/v1/apostas/{key}/grupos",
            headers=auth,
            json={"grupo_ids": [str(detailed.grupos[0])]},
        )
    ).status_code == 200
    assert (
        await client.post(
            f"/api/v1/apostas/{key}/resultado",
            headers=auth,
            json={"estado": "GREEN", "retorno_centavos": 20000},
        )
    ).status_code == 200
    before = await read(api, detailed, "/grupos")
    assert before["total_periodo"]["total_apostas"] == 11
    assert (await client.delete(f"/api/v1/apostas/{key}", headers=auth)).status_code == 200
    deleted = await read(api, detailed, "/diario")
    assert deleted["total_periodo"]["total_apostas"] == 10
    archived = await read(
        api, detailed, "/grupos", {"visibilidade": "apagadas", "grupo_id": str(detailed.grupos[0])}
    )
    assert archived["total_periodo"]["total_apostas"] == 1
    assert archived["grupos"][0]["id"] == str(detailed.grupos[0])
    assert (await client.post(f"/api/v1/apostas/{key}/restaurar", headers=auth)).status_code == 200
    restored = await read(api, detailed, "/grupos")
    assert restored["grupos"] == before["grupos"]
    assert restored["total_periodo"] == before["total_periodo"]
