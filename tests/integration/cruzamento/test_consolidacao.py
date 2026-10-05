"""Synthetic product acceptance on PostgreSQL, using actual intake and workers."""

import asyncio
import copy
import os
from datetime import UTC, datetime, timedelta
from uuid import uuid4

import pytest
from sqlalchemy import func, select, text, update
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.exc import DBAPIError
from sqlalchemy.ext.asyncio import AsyncSession

from bancaemdia import models
from bancaemdia.api.v1.coleta import registrar
from bancaemdia.cli.replay import reconstruir_usuario
from bancaemdia.domain.account_attribution import ResolutionStatus
from bancaemdia.domain.account_attribution_service import (
    InvalidAccountReferenceError,
    attribute_account,
)
from bancaemdia.domain.consolidacao_aposta import ConsolidacaoRecusadaError, consolidate, unlink
from bancaemdia.repositories.aposta_consolidacao import financial_predicate
from bancaemdia.repositories.aposta_repo import ApostaRepo
from bancaemdia.repositories.extrato_repo import ExtratoRepo
from bancaemdia.workers.materialization import ExtracaoDoJson, gravar_coletas, gravar_leitura

pytestmark = pytest.mark.xdist_group("postgres")
PLACEMENT = datetime(2026, 9, 20, 12, tzinfo=UTC)
GAME = datetime(2026, 9, 22, 22, tzinfo=UTC)


@pytest.fixture(scope="module", autouse=True)
def require_postgres():
    if not os.environ.get("TEST_DATABASE_URL"):
        from testcontainers.core.docker_client import DockerClient

        assert DockerClient().client.ping(), "Consolidation acceptance requires real PostgreSQL"


async def owner(session, user):
    await session.execute(
        text("SELECT set_config('app.current_user_id',:u,true)"), {"u": str(user)}
    )


async def accounts(engine_admin, engine_app, user, *, count=1):
    async with engine_admin.begin() as conn:
        await conn.execute(
            insert(models.Casa)
            .values(nome="Betano", dominio="betano.bet.br")
            .on_conflict_do_nothing()
        )
        house = await conn.scalar(select(models.Casa.id).where(models.Casa.nome == "Betano"))
    async with AsyncSession(engine_app, expire_on_commit=False) as session, session.begin():
        await owner(session, user)
        rows = []
        for index in range(count):
            account = models.ContaCasa(
                usuario_id=user,
                casa_id=house,
                apelido=f"synthetic-{index}",
                desde=PLACEMENT - timedelta(days=1),
            )
            session.add(account)
            await session.flush()
            rows.append(account.id)
        # Multiple account identities alone are not a default assignment. Usage
        # windows live in the authoritative table introduced by the reviewed #94.
        if count == 1:
            from bancaemdia.repositories.uso_conta_casa_repo import UsoContaCasaRepo

            await UsoContaCasaRepo().open(
                session, user, house, rows[0], PLACEMENT - timedelta(days=1)
            )
        return house, rows


def house_payload(ticket, *, state=None, return_value=200):
    payload = {
        "id": ticket,
        "totalAmount": 100,
        "totalOdds": 2.0,
        "placedAt": int(PLACEMENT.timestamp() * 1000),
        "bonusType": 0,
        "totalAmountWithCurrency": {"currencyCode": "BRL"},
        "legs": [
            {
                "legItems": [
                    {
                        "eventId": "synthetic-42",
                        "eventName": "Azul - Verde",
                        "startTime": int(GAME.timestamp() * 1000),
                        "selections": [
                            {
                                "description": "Mais de 2.5",
                                "market": "Total de gols",
                                "handicap": 2.5,
                                "odds": 2.0,
                            }
                        ],
                    }
                ]
            }
        ],
    }
    if state:
        payload.update(
            finalBetResult=state,
            finalWinnings=return_value,
            settledAt=int((GAME + timedelta(hours=2)).timestamp() * 1000),
        )
    return payload


def telegram_payload(ticket, *, identity=True, chat=None, message=None):
    return {
        "chat_id": chat or uuid4().int % (2**60),
        "message_id": message or uuid4().int % (2**60),
        "postada_em": PLACEMENT.isoformat(),
        "bilhete": {
            "casa": "Betano",
            "tipo": "simples",
            "evento": "Azul - Verde",
            "odd_total": 2.0,
            "quando": GAME.isoformat(),
            "confianca": 0.99,
            "identidade_bilhete": ticket if identity else None,
            "ocorrido_em": PLACEMENT.isoformat(),
            "stake_unidades": 1.0,
            "selecoes": [
                {
                    "evento": "Azul - Verde",
                    "mercado": "Total de gols",
                    "escolha": "Mais de 2.5",
                    "linha": 2.5,
                    "odd": 2.0,
                }
            ],
        },
    }


async def intake_house(engine, user, house_id, payload):
    async with AsyncSession(engine) as session, session.begin():
        await owner(session, user)
        result, jobs = await registrar(session, user, "betano", house_id, [payload])
        assert result.recusadas == []
    await gravar_coletas(engine, user, jobs)
    async with AsyncSession(engine) as session:
        await owner(session, user)
        return await session.scalar(
            select(models.Aposta.id).where(
                models.Aposta.usuario_id == user, models.Aposta.chave == f"c:betano:{payload['id']}"
            )
        )


async def intake_telegram(engine, user, payload):
    await gravar_leitura(
        engine, user, ExtracaoDoJson.model_validate(payload), payload, "synthetic-media"
    )
    async with AsyncSession(engine) as session:
        await owner(session, user)
        return await session.scalar(
            select(models.Aposta.id).where(
                models.Aposta.usuario_id == user,
                models.Aposta.chat_id == payload["chat_id"],
                models.Aposta.message_id == payload["message_id"],
            )
        )


async def financial(engine, user):
    async with AsyncSession(engine) as session:
        await owner(session, user)
        bets = list(
            await session.scalars(
                select(models.Aposta).where(
                    models.Aposta.usuario_id == user,
                    models.Aposta.selecionada,
                    financial_predicate(),
                )
            )
        )
        relations = list(
            await session.scalars(
                select(models.ApostaConsolidacao)
                .where(models.ApostaConsolidacao.usuario_id == user)
                .order_by(models.ApostaConsolidacao.id)
            )
        )
        return {
            "count": len(bets),
            "stake": sum(b.stake_centavos for b in bets),
            "return": sum(b.retorno_centavos or 0 for b in bets),
            "exposure": sum(b.stake_centavos for b in bets if b.estado == "PENDENTE"),
        }, relations


@pytest.mark.parametrize("order", ["house-first", "telegram-first"])
async def test_real_intake_both_orders_one_fact_settlement_and_source_preservation(
    engine_app, engine_admin, novo_usuario, order
):
    user = await novo_usuario()
    house_id, account_ids = await accounts(engine_admin, engine_app, user)
    ticket = "synthetic-" + uuid4().hex
    raw, tg = house_payload(ticket), telegram_payload(ticket)
    if order == "house-first":
        house = await intake_house(engine_app, user, house_id, raw)
        tip = await intake_telegram(engine_app, user, tg)
    else:
        tip = await intake_telegram(engine_app, user, tg)
        house = await intake_house(engine_app, user, house_id, raw)
    totals, relations = await financial(engine_app, user)
    assert totals == {"count": 1, "stake": 10000, "return": 0, "exposure": 10000}
    assert len(relations) == 1 and relations[0].estado == "active"
    assert relations[0].casa_aposta_id == house and relations[0].telegram_aposta_id == tip
    assert relations[0].conta_casa_id == account_ids[0]
    assert relations[0].contexto["midia_hash"] == "synthetic-media"
    assert relations[0].evidencia["veredicto"]["classification"] == "exact"
    await intake_house(engine_app, user, house_id, house_payload(ticket, state="Win"))
    # Telegram rereading may contradict money; it cannot overwrite the Casa fact or re-enable it.
    updated_tg = copy.deepcopy(tg)
    updated_tg["bilhete"]["odd_total"] = 3.0
    await intake_telegram(engine_app, user, updated_tg)
    totals, repeated = await financial(engine_app, user)
    assert totals == {"count": 1, "stake": 10000, "return": 20000, "exposure": 0}
    assert repeated[0].id == relations[0].id
    async with AsyncSession(engine_app) as session:
        await owner(session, user)
        rows = list(
            await session.scalars(select(models.Aposta).where(models.Aposta.usuario_id == user))
        )
        assert len(rows) == 2 and all(b.selecionada for b in rows)
        assert next(b for b in rows if b.id == house).odd == pytest.approx(2.0)
        assert next(b for b in rows if b.id == tip).odd == pytest.approx(3.0)
        assert (
            await session.scalar(
                select(func.count())
                .select_from(models.ColetaCasa)
                .where(models.ColetaCasa.usuario_id == user)
            )
            == 1
        )
        assert (
            await session.scalar(
                select(func.count())
                .select_from(models.Evento)
                .where(
                    models.Evento.usuario_id == user, models.Evento.tipo == "APOSTAS_CONSOLIDADAS"
                )
            )
            == 2
        )
        ledger, total = await ExtratoRepo().list_page(session, user, {})
        assert total == 1 and ledger[0].resultado_liquido_centavos == 10000
        listed, count = await ApostaRepo().list_page(session, user, {})
        assert count == 1 and listed[0].id == house
    result = await reconstruir_usuario(user, engine=engine_app)
    assert result.apostas == 2
    assert (await financial(engine_app, user))[0] == totals


async def test_default_account_follows_game_day_while_explicit_multiaccount_keeps_actual_account(
    engine_app, engine_admin, novo_usuario
):
    user = await novo_usuario()
    house, ids = await accounts(engine_admin, engine_app, user, count=2)
    switch = PLACEMENT + timedelta(days=1)
    async with AsyncSession(engine_app) as session, session.begin():
        await owner(session, user)
        await session.execute(
            update(models.ContaCasa)
            .where(models.ContaCasa.id == ids[0])
            .values(ate=switch, ativa=False)
        )
        await session.execute(
            update(models.ContaCasa).where(models.ContaCasa.id == ids[1]).values(desde=switch)
        )
        from bancaemdia.repositories.uso_conta_casa_repo import UsoContaCasaRepo

        usage = await UsoContaCasaRepo().open(
            session, user, house, ids[0], PLACEMENT - timedelta(days=1)
        )
        await UsoContaCasaRepo().close(session, usage.id, switch)
        await UsoContaCasaRepo().open(session, user, house, ids[1], switch)

        default = await attribute_account(session, user, "Betano", GAME)
        actual = await attribute_account(session, user, "Betano", GAME, ids[0])
        assert default.conta_casa_id == ids[1]
        assert actual.conta_casa_id == ids[0]
        assert (await attribute_account(session, user, "Betano", switch)).conta_casa_id == ids[1]
        assert (
            await attribute_account(session, user, "Betano", None)
        ).status == ResolutionStatus.NONE
    ticket = "synthetic-" + uuid4().hex
    house_bet = await intake_house(engine_app, user, house, house_payload(ticket))
    await intake_telegram(engine_app, user, telegram_payload(ticket))
    totals, relations = await financial(engine_app, user)
    assert totals["count"] == 1 and relations[0].conta_casa_id == ids[1]
    async with AsyncSession(engine_app) as session:
        await owner(session, user)
        assert (await session.get(models.Aposta, house_bet)).conta_casa_id == ids[1]
    await reconstruir_usuario(user, engine=engine_app)
    assert (await financial(engine_app, user))[1][0].conta_casa_id == ids[1]


@pytest.mark.parametrize("count", [0, 2])
async def test_missing_or_ambiguous_account_never_consolidates_or_chooses_first(
    engine_app, engine_admin, novo_usuario, count
):
    user = await novo_usuario()
    house, _ = await accounts(engine_admin, engine_app, user, count=count)
    ticket = "synthetic-" + uuid4().hex
    await intake_house(engine_app, user, house, house_payload(ticket))
    await intake_telegram(engine_app, user, telegram_payload(ticket))
    totals, relations = await financial(engine_app, user)
    assert totals["stake"] == 20000 and relations == []
    async with AsyncSession(engine_app) as session:
        await owner(session, user)
        assert all(
            b.conta_casa_id is None
            for b in await session.scalars(
                select(models.Aposta).where(models.Aposta.usuario_id == user)
            )
        )
        assert (
            await session.scalar(
                select(func.count())
                .select_from(models.RevisaoPendente)
                .where(
                    models.RevisaoPendente.usuario_id == user,
                    models.RevisaoPendente.resolvido_em.is_(None),
                )
            )
            >= 1
        )


async def test_probable_missing_identity_preserves_two_facts_until_explicit_review(
    engine_app, engine_admin, novo_usuario
):
    user = await novo_usuario()
    house, _ = await accounts(engine_admin, engine_app, user)
    ticket = "synthetic-" + uuid4().hex
    house_bet = await intake_house(engine_app, user, house, house_payload(ticket))
    tip = await intake_telegram(engine_app, user, telegram_payload(ticket, identity=False))
    assert (await financial(engine_app, user))[0]["count"] == 2
    async with AsyncSession(engine_app) as session, session.begin():
        await owner(session, user)
        relation = await consolidate(
            session, user, house_bet, tip, decision="reviewed", actor_id=user
        )
        ident = relation.id
    assert (await financial(engine_app, user))[0]["count"] == 1
    async with AsyncSession(engine_app) as session, session.begin():
        await owner(session, user)
        await unlink(session, user, ident, actor_id=user, reason="synthetic reclassification")
    assert (await financial(engine_app, user))[0]["count"] == 2
    with pytest.raises(ConsolidacaoRecusadaError, match="revisada"):
        async with AsyncSession(engine_app) as session, session.begin():
            await owner(session, user)
            await consolidate(session, user, house_bet, tip)
    async with AsyncSession(engine_app) as session, session.begin():
        await owner(session, user)
        newer = await consolidate(session, user, house_bet, tip, decision="reviewed", actor_id=user)
        assert newer.id != ident
    _, relations = await financial(engine_app, user)
    assert [r.estado for r in relations] == ["unlinked", "active"]


async def test_concurrent_same_pair_retry_and_transaction_rollback(
    engine_app, engine_admin, novo_usuario
):
    user = await novo_usuario()
    house, _ = await accounts(engine_admin, engine_app, user)
    ticket = "synthetic-" + uuid4().hex
    house_bet = await intake_house(engine_app, user, house, house_payload(ticket))
    tip = await intake_telegram(engine_app, user, telegram_payload(ticket, identity=False))
    with pytest.raises(RuntimeError, match="injected"):
        async with AsyncSession(engine_app) as session, session.begin():
            await owner(session, user)
            await consolidate(session, user, house_bet, tip, decision="reviewed", actor_id=user)
            raise RuntimeError("injected rollback after relation/events/source update")
    assert (await financial(engine_app, user))[1] == []

    async def apply():
        async with AsyncSession(engine_app) as session, session.begin():
            await owner(session, user)
            relation = await consolidate(
                session, user, house_bet, tip, decision="reviewed", actor_id=user
            )
            return relation.id

    ids = await asyncio.wait_for(asyncio.gather(*(apply() for _ in range(10))), timeout=30)
    assert len(set(ids)) == 1
    assert (await financial(engine_app, user))[0]["stake"] == 10000
    async with AsyncSession(engine_app) as session:
        await owner(session, user)
        assert (
            await session.scalar(
                select(func.count())
                .select_from(models.Evento)
                .where(
                    models.Evento.usuario_id == user, models.Evento.tipo == "APOSTAS_CONSOLIDADAS"
                )
            )
            == 2
        )


async def test_cross_tenant_rls_fk_and_explicit_reference_fail_closed(
    engine_app, engine_admin, novo_usuario
):
    user, foreign = await novo_usuario(), await novo_usuario()
    house, ids = await accounts(engine_admin, engine_app, user)
    _, foreign_ids = await accounts(engine_admin, engine_app, foreign)
    ticket = "synthetic-" + uuid4().hex
    casa = await intake_house(engine_app, user, house, house_payload(ticket))
    tip = await intake_telegram(engine_app, user, telegram_payload(ticket))
    foreign_tip = await intake_telegram(engine_app, foreign, telegram_payload(ticket))
    relation = (await financial(engine_app, user))[1][0]
    async with AsyncSession(engine_app) as session, session.begin():
        await owner(session, foreign)
        assert await session.get(models.ApostaConsolidacao, relation.id) is None
        assert (
            await session.execute(
                update(models.ApostaConsolidacao)
                .where(models.ApostaConsolidacao.id == relation.id)
                .values(estado="unlinked", desvinculada_em=GAME)
            )
        ).rowcount == 0
    with pytest.raises(InvalidAccountReferenceError):
        async with AsyncSession(engine_app) as session, session.begin():
            await owner(session, user)
            await attribute_account(session, user, "Betano", GAME, foreign_ids[0])
    with pytest.raises(DBAPIError):
        async with AsyncSession(engine_app) as session, session.begin():
            await owner(session, user)
            session.add(
                models.ApostaConsolidacao(
                    usuario_id=user,
                    casa_aposta_id=casa,
                    telegram_aposta_id=foreign_tip,
                    conta_casa_id=ids[0],
                    estado="active",
                    decisao="reviewed",
                    versao="synthetic",
                    evidencia={},
                    contexto={},
                    ator="synthetic",
                )
            )
            await session.flush()
    with pytest.raises(ConsolidacaoRecusadaError):
        async with AsyncSession(engine_app) as session, session.begin():
            await owner(session, foreign)
            await consolidate(session, foreign, casa, tip, decision="reviewed", actor_id=foreign)


def application(engine, user):
    from types import SimpleNamespace

    from fastapi import FastAPI

    from bancaemdia.api.deps import get_current_user, get_current_user_snapshot
    from bancaemdia.api.v1 import apostas, revisao
    from bancaemdia.db.session import get_db, get_db_snapshot

    app = FastAPI()
    app.include_router(apostas.router)
    app.include_router(revisao.router)

    async def database():
        async with AsyncSession(engine, expire_on_commit=False) as session:
            await owner(session, user)
            yield session

    def identity():
        return SimpleNamespace(id=user)

    for dep in (get_current_user, get_current_user_snapshot):
        app.dependency_overrides[dep] = identity
    for dep in (get_db, get_db_snapshot):
        app.dependency_overrides[dep] = database
    return app


@pytest.mark.parametrize(
    ("state", "paid", "expected"),
    [
        ("Win", 200, "GREEN"),
        ("Void", 100, "ANULADA"),
        ("Cashout", 75, "CASHOUT"),
        ("Lose", 0, "RED"),
    ],
)
async def test_lifecycle_keeps_one_fact_with_independent_paid_value(
    engine_app, engine_admin, novo_usuario, state, paid, expected
):
    user = await novo_usuario()
    house, ids = await accounts(engine_admin, engine_app, user)
    ticket = uuid4().hex
    await intake_telegram(engine_app, user, telegram_payload(ticket))
    casa = await intake_house(engine_app, user, house, house_payload(ticket))
    await intake_house(
        engine_app, user, house, house_payload(ticket, state=state, return_value=paid)
    )
    totals, relations = await financial(engine_app, user)
    assert totals == {"count": 1, "stake": 10000, "return": paid * 100, "exposure": 0}
    assert len(relations) == 1
    async with AsyncSession(engine_app) as session:
        await owner(session, user)
        assert (await session.get(models.Aposta, casa)).estado == expected
        from bancaemdia.repositories.movimento_repo import MovimentoRepo

        balances = await MovimentoRepo().aggregate_saldos_by_usuario(session, user)
        assert balances[ids[0]].apostado_centavos == 10000
        assert balances[ids[0]].retornado_centavos == paid * 100
        from bancaemdia.repositories.account_financials_repo import AccountFinancialsRepo

        metrics = await AccountFinancialsRepo().bets(
            session, user, desde=None, ate=None, casa_id=None, titular_id=None, conta_casa_id=ids[0]
        )
        assert len(metrics) == 1 and metrics[0]["apostas"] == 1
        assert metrics[0]["retorno_centavos"] == (0 if expected == "ANULADA" else paid * 100)
    await reconstruir_usuario(user, engine=engine_app)
    assert (await financial(engine_app, user))[0] == totals


@pytest.mark.parametrize("action", ["MESMA", "CORRIGIR"])
async def test_actual_review_api_confirms_or_durably_rejects_the_candidate(
    engine_app, engine_admin, novo_usuario, action
):
    import httpx

    user = await novo_usuario()
    house, _ = await accounts(engine_admin, engine_app, user)
    ticket = uuid4().hex
    await intake_house(engine_app, user, house, house_payload(ticket))
    await intake_telegram(engine_app, user, telegram_payload(ticket, identity=False))
    async with AsyncSession(engine_app) as session:
        await owner(session, user)
        pair = await session.scalar(
            select(models.CruzamentoCandidato).where(models.CruzamentoCandidato.usuario_id == user)
        )
        assert pair.status == "probable" and pair.revisao_id
        review_id = pair.revisao_id
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=application(engine_app, user)), base_url="http://test"
    ) as client:
        request = {"acao": action}
        if action == "CORRIGIR":
            request["aposta_corrigida"] = {"odd": 2.1}
        response = await client.post(f"/api/v1/revisao/{review_id}/resolver", json=request)
        assert response.status_code == 200, response.text
    totals, relations = await financial(engine_app, user)
    assert totals["count"] == (1 if action == "MESMA" else 2)
    assert len(relations) == 1 and relations[0].estado == (
        "active" if action == "MESMA" else "rejected"
    )
    if action == "CORRIGIR":
        async with AsyncSession(engine_app) as session, session.begin():
            await owner(session, user)
            with pytest.raises(ConsolidacaoRecusadaError, match="nova decisão revisada"):
                await consolidate(
                    session, user, relations[0].casa_aposta_id, relations[0].telegram_aposta_id
                )
    await reconstruir_usuario(user, engine=engine_app)
    assert (await financial(engine_app, user))[0] == totals


async def test_consolidation_diagnostics_keep_sensitive_sentinels_out_of_labels_and_spans(
    engine_app, engine_admin, novo_usuario, monkeypatch, caplog
):
    from opentelemetry.sdk.trace import TracerProvider
    from opentelemetry.sdk.trace.export import SimpleSpanProcessor
    from opentelemetry.sdk.trace.export.in_memory_span_exporter import InMemorySpanExporter

    from bancaemdia.domain.consolidacao_aposta import outcomes
    from bancaemdia.observability import tracing

    sentinel = "synthetic-secret-signed-url-token@example.invalid"
    exporter = InMemorySpanExporter()
    provider = TracerProvider()
    provider.add_span_processor(SimpleSpanProcessor(exporter))
    monkeypatch.setattr(tracing, "_tracer", lambda: provider.get_tracer("synthetic-consolidation"))
    user = await novo_usuario()
    house, _ = await accounts(engine_admin, engine_app, user)
    tg = telegram_payload(sentinel)
    tg["bilhete"]["tipster"] = sentinel
    try:
        await intake_house(engine_app, user, house, house_payload(sentinel))
        await intake_telegram(engine_app, user, tg)
        assert (await financial(engine_app, user))[0]["count"] == 1
        labels = [sample.labels for family in outcomes.collect() for sample in family.samples]
        assert all(set(item) <= {"decision", "result"} for item in labels)
        assert sentinel not in repr(labels)
        assert sentinel not in caplog.text
        assert sentinel not in repr([
            (span.attributes, span.events, span.status.description)
            for span in exporter.get_finished_spans()
        ])
    finally:
        provider.shutdown()


async def test_source_edit_delete_restore_and_private_detail_do_not_bypass_relation(
    engine_app, engine_admin, novo_usuario
):
    import httpx

    user, foreign = await novo_usuario(), await novo_usuario()
    house, _ = await accounts(engine_admin, engine_app, user)
    ticket = uuid4().hex
    casa = await intake_house(engine_app, user, house, house_payload(ticket))
    tip = await intake_telegram(engine_app, user, telegram_payload(ticket))
    async with AsyncSession(engine_app) as session:
        await owner(session, user)
        sources = {
            bet.id: bet.chave
            for bet in await session.scalars(
                select(models.Aposta).where(models.Aposta.usuario_id == user)
            )
        }
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=application(engine_app, user)), base_url="http://test"
    ) as client:
        response = await client.get("/api/v1/apostas/" + sources[tip])
        assert response.status_code == 200
        assert response.json()["fonte_contextual"] is True
        assert response.json()["aposta"]["apagada"] is False
        for ident, expected_count in ((tip, 1), (casa, 0)):
            url = "/api/v1/apostas/" + sources[ident]
            assert (await client.delete(url)).status_code == 200
            assert (await financial(engine_app, user))[0]["count"] == expected_count
            assert (await client.post(url + "/restaurar")).status_code == 200
            assert (await financial(engine_app, user))[0]["count"] == 1
        assert (
            await client.patch(
                "/api/v1/apostas/" + sources[tip], json={"odd": 3.5, "stake_unidades": 5}
            )
        ).status_code == 200
        assert (await financial(engine_app, user))[0]["stake"] == 10000
        assert (
            await client.patch("/api/v1/apostas/" + sources[casa], json={"stake_unidades": 2})
        ).status_code == 200
        assert (await financial(engine_app, user))[0]["stake"] == 20000
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=application(engine_app, foreign)), base_url="http://test"
    ) as client:
        assert (await client.get("/api/v1/apostas/" + sources[tip])).status_code == 404
        relation = (await financial(engine_app, user))[1][0]
        assert (
            await client.post(
                f"/api/v1/consolidacoes/{relation.id}/desvincular", json={"motivo": "foreign"}
            )
        ).status_code == 409
    await reconstruir_usuario(user, engine=engine_app)
    assert (await financial(engine_app, user))[0]["stake"] == 20000


@pytest.mark.parametrize("side", ["casa", "telegram"])
async def test_symmetric_competition_serializes_reviewed_decisions_without_deadlock(
    engine_app, engine_admin, novo_usuario, side
):
    user = await novo_usuario()
    house, _ = await accounts(engine_admin, engine_app, user)
    first, second = uuid4().hex, uuid4().hex
    houses = [await intake_house(engine_app, user, house, house_payload(first))]
    tips = [await intake_telegram(engine_app, user, telegram_payload(first, identity=False))]
    if side == "casa":
        houses.append(await intake_house(engine_app, user, house, house_payload(second)))
        pairs = [(c, tips[0]) for c in houses]
    else:
        tips.append(
            await intake_telegram(engine_app, user, telegram_payload(first, identity=False))
        )
        pairs = [(houses[0], t) for t in tips]
    assert (await financial(engine_app, user))[1] == []

    async def apply(c, t):
        try:
            async with AsyncSession(engine_app) as session, session.begin():
                await owner(session, user)
                result = await consolidate(session, user, c, t, decision="reviewed", actor_id=user)
                return result.id
        except ConsolidacaoRecusadaError:
            return None

    results = await asyncio.wait_for(asyncio.gather(*(apply(c, t) for c, t in pairs)), timeout=20)
    assert sum(result is not None for result in results) == 1
    totals, relations = await financial(engine_app, user)
    assert totals["count"] == 2 and totals["stake"] == 20000
    assert len(relations) == 1


async def test_context_conflict_requires_choice_and_preserves_both_original_values(
    engine_app, engine_admin, novo_usuario
):
    user = await novo_usuario()
    house, _ = await accounts(engine_admin, engine_app, user)
    ticket = uuid4().hex
    casa = await intake_house(engine_app, user, house, house_payload(ticket))
    async with AsyncSession(engine_app) as session, session.begin():
        await owner(session, user)
        session.add(
            models.Evento(
                usuario_id=user,
                aposta_chave=f"c:betano:{ticket}",
                tipo="CORRECAO_MANUAL",
                fonte="manual",
                payload_json={"tipster": "source-house"},
            )
        )
    tg = telegram_payload(ticket)
    tg["bilhete"]["tipster"] = "source-telegram"
    tip = await intake_telegram(engine_app, user, tg)
    assert (await financial(engine_app, user))[1] == []
    async with AsyncSession(engine_app) as session, session.begin():
        await owner(session, user)
        result = await consolidate(
            session, user, casa, tip, decision="reviewed", actor_id=user, tipster_choice="telegram"
        )
        assert result.contexto["tipster_casa"] == "source-house"
        assert result.contexto["tipster_telegram"] == "source-telegram"
        assert result.contexto["escolha"] == "telegram"
    assert (await financial(engine_app, user))[0]["stake"] == 10000


async def test_manual_account_game_date_and_explicit_reference_survive_api_and_replay(
    engine_app, engine_admin, novo_usuario
):
    import httpx

    user = await novo_usuario()
    house, ids = await accounts(engine_admin, engine_app, user, count=2)
    boundary = PLACEMENT + timedelta(days=1)
    async with AsyncSession(engine_app) as session, session.begin():
        await owner(session, user)
        await session.execute(
            update(models.ContaCasa)
            .where(models.ContaCasa.id == ids[0])
            .values(ate=boundary, ativa=False)
        )
        await session.execute(
            update(models.ContaCasa).where(models.ContaCasa.id == ids[1]).values(desde=boundary)
        )
        from bancaemdia.repositories.uso_conta_casa_repo import UsoContaCasaRepo

        usage = await UsoContaCasaRepo().open(
            session, user, house, ids[0], PLACEMENT - timedelta(days=1)
        )
        await UsoContaCasaRepo().close(session, usage.id, boundary)
        await UsoContaCasaRepo().open(session, user, house, ids[1], boundary)

    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=application(engine_app, user)), base_url="http://test"
    ) as client:
        body = {
            "casa": "Betano",
            "odd": 2,
            "stake_unidades": 1,
            "data_aposta": PLACEMENT.isoformat(),
            "data_jogo": GAME.isoformat(),
        }
        for reference, expected in ((None, ids[1]), (ids[0], ids[0])):
            request = {**body, **({"conta_casa_ref": reference} if reference else {})}
            response = await client.post("/api/v1/apostas", json=request)
            assert response.status_code == 201, response.text
            assert response.json()["aposta"]["conta_casa_id"] == expected
        response = await client.post("/api/v1/apostas", json={**body, "data_jogo": None})
        assert response.status_code == 201 and response.json()["aposta"]["conta_casa_id"] is None
        assert (
            await client.post("/api/v1/apostas", json={**body, "conta_casa_ref": 2**63 - 1})
        ).status_code == 422
    await reconstruir_usuario(user, engine=engine_app, dry_run=True)
    await reconstruir_usuario(user, engine=engine_app)
    async with AsyncSession(engine_app) as session:
        await owner(session, user)
        rows = list(
            await session.scalars(
                select(models.Aposta)
                .where(models.Aposta.usuario_id == user)
                .order_by(models.Aposta.id)
            )
        )
        assert [r.conta_casa_id for r in rows] == [ids[1], ids[0], None]


@pytest.mark.parametrize("complete", [True, False])
async def test_multiple_ticket_requires_every_leg_for_automatic_consolidation(
    engine_app, engine_admin, novo_usuario, complete
):
    user = await novo_usuario()
    house, _ = await accounts(engine_admin, engine_app, user)
    ticket = uuid4().hex
    raw, tg = house_payload(ticket), telegram_payload(ticket)
    second = copy.deepcopy(raw["legs"][0]["legItems"][0])
    second.update(eventId="synthetic-43", eventName="Laranja - Branco")
    raw["legs"].append({"legItems": [second]})
    tg["bilhete"]["tipo"] = "multipla"
    tg["bilhete"]["selecoes"][0]["odd"] = None
    if complete:
        tg["bilhete"]["selecoes"].append({
            **tg["bilhete"]["selecoes"][0],
            "evento": "Laranja - Branco",
        })
    await intake_house(engine_app, user, house, raw)
    await intake_telegram(engine_app, user, tg)
    totals, relations = await financial(engine_app, user)
    assert totals["count"] == (1 if complete else 2)
    assert len(relations) == (1 if complete else 0)


@pytest.mark.parametrize("change", ["correction", "deleted", "saturated"])
async def test_revalidation_refuses_previous_exact_after_source_or_search_changes(
    engine_app, engine_admin, novo_usuario, change
):
    user = await novo_usuario()
    house, _ = await accounts(engine_admin, engine_app, user, count=0)
    ticket = uuid4().hex
    casa = await intake_house(engine_app, user, house, house_payload(ticket))
    tip = await intake_telegram(engine_app, user, telegram_payload(ticket))
    async with AsyncSession(engine_app) as session:
        await owner(session, user)
        candidate = await session.scalar(
            select(models.CruzamentoCandidato).where(models.CruzamentoCandidato.usuario_id == user)
        )
        assert candidate.status == "exact"
        ident = candidate.id
    if change == "saturated":
        # Each filler crosses the real intake/matching path, including hidden saturation markers.
        for _ in range(201):
            await intake_telegram(engine_app, user, telegram_payload(uuid4().hex))
    else:
        async with AsyncSession(engine_app) as session, session.begin():
            await owner(session, user)
            values = {"odd": 3} if change == "correction" else {"selecionada": False}
            await session.execute(
                update(models.Aposta).where(models.Aposta.id == tip).values(**values)
            )
    await accounts(engine_admin, engine_app, user)
    with pytest.raises(ConsolidacaoRecusadaError):
        async with AsyncSession(engine_app) as session, session.begin():
            await owner(session, user)
            await consolidate(session, user, casa, tip, candidate_id=ident)
    assert (await financial(engine_app, user))[1] == []


async def test_panel_refresh_filters_export_and_cash_use_same_single_fact(
    engine_app, engine_admin, novo_usuario, tmp_path
):
    from openpyxl import load_workbook

    from bancaemdia.cli.refresh_painel import refresh_painel
    from bancaemdia.domain.painel import FiltrosPainel, SecaoExportacao
    from bancaemdia.exportacao.painel_xlsx import AbaXlsxAssincrona, gerar_painel_xlsx_assincrono
    from bancaemdia.repositories.movimento_repo import MovimentoRepo
    from bancaemdia.repositories.painel_repo import PainelRepo

    user = await novo_usuario()
    house, ids = await accounts(engine_admin, engine_app, user)
    context_name = "synthetic-panel-" + uuid4().hex
    async with engine_admin.begin() as conn:
        await conn.execute(insert(models.Tipster).values(nome=context_name))
    ticket = uuid4().hex
    await intake_house(engine_app, user, house, house_payload(ticket))
    tg = telegram_payload(ticket)
    tg["bilhete"]["tipster"] = context_name
    await intake_telegram(engine_app, user, tg)
    await intake_house(engine_app, user, house, house_payload(ticket, state="Win"))
    async with AsyncSession(engine_app) as session, session.begin():
        await owner(session, user)
        session.add(
            models.Movimento(
                usuario_id=user,
                conta_casa_id=ids[0],
                tipo="DEPOSITO",
                valor_centavos=50000,
                ocorrido_em=PLACEMENT - timedelta(days=1),
            )
        )
    refreshed = await refresh_painel(engine_admin)
    async with AsyncSession(engine_app) as session:
        await owner(session, user)
        balances = await MovimentoRepo().aggregate_saldos_by_usuario(session, user)
        assert balances[ids[0]].saldo_centavos == 60000
        contextual_tipster = await session.scalar(
            select(models.Aposta.tipster_id).where(
                models.Aposta.usuario_id == user, models.Aposta.origem == "casa"
            )
        )
        assert contextual_tipster is not None
        for filter_house, filter_tipster in (
            (None, None),
            (house, None),
            (house, contextual_tipster),
        ):
            filters = FiltrosPainel.criar(
                "all",
                casa_id=filter_house,
                tipster_id=filter_tipster,
                agora=GAME + timedelta(days=1),
            )
            panel = await PainelRepo().consultar(session, user, filters)
            assert panel.resumo.total_apostas == 1
            assert panel.resumo.giro_centavos == 10000
            assert panel.resumo.lucro_centavos == 10000
            assert panel.resumo.roi_basis_points == 10000
            assert panel.frescor.atualizado_em == refreshed
            rows = [
                r
                async for r in PainelRepo().iterar_exportacao(
                    session, user, filters, secoes=(SecaoExportacao.RESUMO,)
                )
            ]
            assert len(rows) == 1 and rows[0].valores["giro_centavos"] == 10000

            async def excel_rows(source_rows=rows):
                for row in source_rows:
                    await asyncio.sleep(0)
                    yield [row.valores["giro_centavos"], row.valores["lucro_centavos"]]

            file = await gerar_painel_xlsx_assincrono(
                [AbaXlsxAssincrona("Resumo", ("Stake", "Lucro"), excel_rows())],
                diretorio_temporario=tmp_path,
            )
            workbook = load_workbook(file.caminho, read_only=True, data_only=True)
            assert list(workbook.active.values) == [("Stake", "Lucro"), (10000, 10000)]
            workbook.close()


async def test_replay_reconstructs_missing_relation_from_audit_without_live_candidates(
    engine_app, engine_admin, novo_usuario
):
    from bancaemdia.services.consolidacao_replay import replay_relations

    user = await novo_usuario()
    house, _ = await accounts(engine_admin, engine_app, user)
    ticket = uuid4().hex
    await intake_house(engine_app, user, house, house_payload(ticket))
    await intake_telegram(engine_app, user, telegram_payload(ticket))
    totals, original = await financial(engine_app, user)
    async with engine_admin.begin() as conn:
        # Fault injection only on synthetic projections; never a product deletion capability.
        await conn.execute(
            text("ALTER TABLE aposta_consolidacoes DISABLE TRIGGER consolidation_immutable")
        )
        await conn.execute(
            text("DELETE FROM aposta_consolidacoes WHERE usuario_id=:u"), {"u": user}
        )
        await conn.execute(
            text("ALTER TABLE aposta_consolidacoes ENABLE TRIGGER consolidation_immutable")
        )
    async with AsyncSession(engine_app) as session, session.begin():
        await owner(session, user)
        assert await replay_relations(session, user) == 1
    restored = (await financial(engine_app, user))[1]
    assert restored[0].id == original[0].id
    assert restored[0].evidencia == original[0].evidencia
    assert restored[0].contexto == original[0].contexto
    await reconstruir_usuario(user, engine=engine_app)
    assert (await financial(engine_app, user))[0] == totals


async def test_concurrent_actual_intakes_converge_without_deadlock(
    engine_app, engine_admin, novo_usuario
):
    user = await novo_usuario()
    house, _ = await accounts(engine_admin, engine_app, user)
    ticket = uuid4().hex
    await asyncio.wait_for(
        asyncio.gather(
            intake_house(engine_app, user, house, house_payload(ticket)),
            intake_telegram(engine_app, user, telegram_payload(ticket)),
        ),
        timeout=20,
    )
    totals, relations = await financial(engine_app, user)
    assert totals == {"count": 1, "stake": 10000, "return": 0, "exposure": 10000}
    assert len(relations) == 1


async def test_unlink_preserves_later_manual_context_and_user_deletion(
    engine_app, engine_admin, novo_usuario
):
    import httpx

    user = await novo_usuario()
    house, _ = await accounts(engine_admin, engine_app, user)
    ticket = uuid4().hex
    casa = await intake_house(engine_app, user, house, house_payload(ticket))
    tip = await intake_telegram(engine_app, user, telegram_payload(ticket))
    relation = (await financial(engine_app, user))[1][0]
    async with engine_admin.begin() as conn:
        context_id = await conn.scalar(
            insert(models.Tipster).values(nome=uuid4().hex).returning(models.Tipster.id)
        )
    async with AsyncSession(engine_app) as session, session.begin():
        await owner(session, user)
        sources = {
            b.id: b.chave
            for b in await session.scalars(
                select(models.Aposta).where(models.Aposta.usuario_id == user)
            )
        }
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=application(engine_app, user)), base_url="http://test"
    ) as client:
        assert (
            await client.patch("/api/v1/apostas/" + sources[casa], json={"tipster_id": context_id})
        ).status_code == 200
        assert (await client.delete("/api/v1/apostas/" + sources[tip])).status_code == 200
        result = await client.post(
            f"/api/v1/consolidacoes/{relation.id}/desvincular",
            json={"motivo": "synthetic independent sources"},
        )
        assert result.status_code == 200, result.text
    await reconstruir_usuario(user, engine=engine_app)
    async with AsyncSession(engine_app) as session:
        await owner(session, user)
        actual = await session.get(models.Aposta, casa)
        deleted = await session.get(models.Aposta, tip)
        assert actual.tipster_id == context_id and not deleted.selecionada
    assert (await financial(engine_app, user))[0]["count"] == 1


async def test_review_media_is_private_and_cross_tenant_ids_cannot_fetch_it(
    engine_app, engine_admin, novo_usuario
):
    import httpx

    user, other = await novo_usuario(), await novo_usuario()
    house, _ = await accounts(engine_admin, engine_app, user)
    ticket = uuid4().hex
    await intake_house(engine_app, user, house, house_payload(ticket))
    await intake_telegram(engine_app, user, telegram_payload(ticket, identity=False))
    async with engine_admin.begin() as conn:
        media_hash = uuid4().hex
        await conn.execute(insert(models.Midia).values(hash=media_hash, tipo="image/png", bytes=16))
        await conn.execute(
            insert(models.MidiaArquivo).values(hash=media_hash, conteudo=b"synthetic-image")
        )
    async with AsyncSession(engine_app) as session, session.begin():
        await owner(session, user)
        review = await session.scalar(
            select(models.RevisaoPendente).where(
                models.RevisaoPendente.usuario_id == user,
                models.RevisaoPendente.motivo == "cruzamento_candidato",
            )
        )
        assert review is not None
        review.midia_hash = media_hash
        ident = review.id
    for tenant, expected in ((user, 200), (other, 404)):
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=application(engine_app, tenant)),
            base_url="http://test",
        ) as client:
            response = await client.get(f"/api/v1/revisao/{ident}/foto")
            assert response.status_code == expected, response.text
            if expected == 200:
                assert response.content == b"synthetic-image"
                assert response.headers["cache-control"] == "private, no-store"


async def test_excel_import_and_update_use_game_date_with_explicit_multiaccount_exception(
    engine_app, engine_admin, novo_usuario
):
    from io import BytesIO

    import httpx
    from openpyxl import Workbook

    user = await novo_usuario()
    house, ids = await accounts(engine_admin, engine_app, user, count=2)
    boundary = PLACEMENT + timedelta(days=1)
    async with AsyncSession(engine_app) as session, session.begin():
        await owner(session, user)
        await session.execute(
            update(models.ContaCasa)
            .where(models.ContaCasa.id == ids[0])
            .values(ate=boundary, ativa=False)
        )
        await session.execute(
            update(models.ContaCasa).where(models.ContaCasa.id == ids[1]).values(desde=boundary)
        )
        from bancaemdia.repositories.uso_conta_casa_repo import UsoContaCasaRepo

        usage = await UsoContaCasaRepo().open(
            session, user, house, ids[0], PLACEMENT - timedelta(days=1)
        )
        await UsoContaCasaRepo().close(session, usage.id, boundary)
        await UsoContaCasaRepo().open(session, user, house, ids[1], boundary)

    def content(revised):
        book = Workbook()
        book.active.append([
            "casa",
            "data_aposta",
            "data_jogo",
            "odd",
            "stake_unidades",
            "atualizada_em",
            "conta_casa_ref",
        ])
        for explicit in (None, ids[0]):
            book.active.append([
                "Betano",
                PLACEMENT.isoformat(),
                GAME.isoformat(),
                2,
                1,
                revised.isoformat(),
                explicit,
            ])
        stream = BytesIO()
        book.save(stream)
        book.close()
        return stream.getvalue()

    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=application(engine_app, user)), base_url="http://test"
    ) as client:
        for index in (0, 1):
            response = await client.post(
                "/api/v1/apostas/importar-planilha",
                data={"origem_id": "synthetic-game-sheet"},
                files={
                    "arquivo": (
                        "test.xlsx",
                        content(GAME + timedelta(days=index)),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert response.status_code == 200, response.text
            assert response.json()["criadas" if index == 0 else "atualizadas"] == 2
    await reconstruir_usuario(user, engine=engine_app)
    async with AsyncSession(engine_app) as session:
        await owner(session, user)
        rows = list(
            await session.scalars(
                select(models.Aposta)
                .where(models.Aposta.usuario_id == user)
                .order_by(models.Aposta.id)
            )
        )
        assert [r.conta_casa_id for r in rows] == [ids[1], ids[0]]
        from bancaemdia.domain.consolidacao_aposta import _state

        states = [await _state(session, user, row.chave) for row in rows]
        assert [state["conta_referencia_explicita"] for state in states] == [False, True]
