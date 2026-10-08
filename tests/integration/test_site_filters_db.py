"""Required site-filter acceptance: actual FastAPI, SQL, application role and RLS.

Only the legacy JWKS network is local; signatures and every database query are real.
The hosted HttpOnly-session journey separately exercises the same production app.
"""

import io
import json
import time
from dataclasses import dataclass
from datetime import UTC, datetime
from uuid import uuid4

import httpx
import jwt
import pytest
from cryptography.hazmat.primitives.asymmetric import rsa
from openpyxl import load_workbook
from sqlalchemy import func, select, text
from sqlalchemy.exc import DBAPIError
from sqlalchemy.ext.asyncio import AsyncSession

from bancaemdia import models
from bancaemdia.auth import jwt as jwt_auth
from bancaemdia.auth import middleware as auth_middleware
from bancaemdia.config import get_settings
from bancaemdia.core.context import current_user_id
from bancaemdia.db.session import get_db, get_db_snapshot
from bancaemdia.main import app
from bancaemdia.repositories.evento_repo import EventoRepo

pytestmark = pytest.mark.xdist_group("postgres")


@dataclass
class Dataset:
    user: int
    other: int
    casas: list[int]
    contas: list[int]
    titulares: list[int]
    bancas: list[int]
    grupos: list[int]
    tipsters: list[int]
    mercados: list[int]
    competicoes: list[int]
    keys: list[str]
    other_ids: dict[str, int]


@pytest.fixture
async def dataset(engine_app, como, novo_usuario):
    user, other = await novo_usuario(), await novo_usuario()
    suffix = uuid4().hex
    async with como(engine_app, user) as session:
        casas = [models.Casa(nome=f"Filtro Casa {i} {suffix}", ativa=i == 0) for i in range(2)]
        tipsters = [models.Tipster(nome=f"Filtro Tipster {i} {suffix}") for i in range(2)]
        mercados = [models.Mercado(nome=f"Filtro Mercado {i} {suffix}") for i in range(2)]
        competicoes = [models.Competicao(nome=f"Filtro Competição {i} {suffix}") for i in range(2)]
        titulares = [
            models.Titular(usuario_id=user, nome=f"Titular {i}", arquivado=i == 1) for i in range(2)
        ]
        bancas = [models.Banca(usuario_id=user, nome=f"Banca {i}") for i in range(2)]
        grupos = [
            models.GrupoAposta(usuario_id=user, nome=f"Grupo {i}", arquivado=i == 1)
            for i in range(2)
        ]
        session.add_all(casas + tipsters + mercados + competicoes + titulares + bancas + grupos)
        await session.flush()
        contas = [
            models.ContaCasa(
                usuario_id=user,
                casa_id=casas[i].id,
                titular_id=titulares[i].id,
                banca_id=bancas[1 - i].id,
                apelido=f"Conta {i}",
                ativa=i == 0,
            )
            for i in range(2)
        ]
        session.add_all(contas)
        await session.flush()
        # Independent, predeclared financial examples: profit +100, -100, 0, +300,
        # +50, -50, +25; a review amount +899 must never enter financial sums.
        estados = [
            "GREEN",
            "RED",
            "PENDENTE",
            "ANULADA",
            "GREEN",
            "MEIO_GREEN",
            "MEIO_RED",
            "CASHOUT",
            "GREEN",
            "GREEN",
            "RED",
            "PENDENTE",
        ]
        stakes = [100, 100, 100, 100, 0, 100, 100, 100, 100, 100, 100, 100]
        retornos = [200, 0, None, 100, 300, 150, 50, 125, 999, None, 0, None]
        rows = []
        for i in range(12):
            key = f"m:{suffix}:{i}"
            # 02:59:59 UTC is previous Sao Paulo civil day; 03:00:00 is inclusive.
            instant = (
                datetime(2026, 9, 10, 2, 59, 59, tzinfo=UTC)
                if i == 0
                else datetime(2026, 9, 10, 3, 0, tzinfo=UTC)
                if i < 10
                else None
            )
            row = models.Aposta(
                usuario_id=user,
                chave=key,
                origem="manual" if i % 2 == 0 else "planilha",
                estado=estados[i],
                stake_unidades=1,
                stake_centavos=stakes[i],
                valor_aposta_centavos=100,
                odd=4 if i == 4 else 2,
                retorno_centavos=retornos[i],
                freebet=i == 4,
                revisao_grave=i == 8,
                selecionada=i < 10,
                conta_casa_id=contas[i % 2].id if i < 10 else None,
                banca_id=bancas[i % 2].id if i < 10 else None,
                tipster_id=tipsters[i % 2].id if i < 10 else None,
                mercado_id=mercados[i % 2].id if i < 10 else None,
                competicao_id=competicoes[i % 2].id if i < 10 else None,
                data_aposta=instant,
                data_jogo=datetime(2026, 10, 1, tzinfo=UTC),
            )
            session.add(row)
            rows.append(row)
        await session.flush()
        session.add_all([
            models.ApostaGrupo(usuario_id=user, aposta_id=row.id, grupo_id=g.id)
            for row in rows[:10]
            for g in grupos
        ])
        await session.commit()
        result = Dataset(
            user,
            other,
            [v.id for v in casas],
            [v.id for v in contas],
            [v.id for v in titulares],
            [v.id for v in bancas],
            [v.id for v in grupos],
            [v.id for v in tipsters],
            [v.id for v in mercados],
            [v.id for v in competicoes],
            [v.chave for v in rows],
            {},
        )
    async with como(engine_app, other) as session:
        bank = models.Banca(usuario_id=other, nome="Banca estrangeira")
        holder = models.Titular(usuario_id=other, nome="Titular estrangeiro")
        group = models.GrupoAposta(usuario_id=other, nome="Grupo estrangeiro")
        session.add_all([bank, holder, group])
        await session.flush()
        account = models.ContaCasa(
            usuario_id=other,
            casa_id=casas[0].id,
            titular_id=holder.id,
            banca_id=bank.id,
            apelido="Conta estrangeira",
        )
        session.add(account)
        await session.flush()
        session.add(
            models.Aposta(
                usuario_id=other,
                chave=f"m:{suffix}:other",
                origem="manual",
                stake_unidades=1,
                stake_centavos=100,
                valor_aposta_centavos=100,
                estado="GREEN",
                retorno_centavos=10_000,
                conta_casa_id=account.id,
                banca_id=bank.id,
            )
        )
        await session.commit()
        result.other_ids = {
            "banca_id": bank.id,
            "titular_id": holder.id,
            "grupo_id": group.id,
            "conta_casa_id": account.id,
        }
    return result


@pytest.fixture
async def api(engine_app, monkeypatch):
    pair = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    public = {
        **json.loads(jwt.algorithms.RSAAlgorithm.to_jwk(pair.public_key())),
        "kid": "site-filter-test",
    }
    cache = jwt_auth.JWKSCache(
        "https://issuer.test/jwks",
        "RS256",
        httpx.MockTransport(lambda request: httpx.Response(200, json={"keys": [public]})),
    )
    monkeypatch.setattr(auth_middleware, "get_jwks_cache", lambda: cache)

    async def db():
        async with AsyncSession(
            engine_app.execution_options(isolation_level="REPEATABLE READ"), expire_on_commit=False
        ) as session:
            uid = current_user_id.get()
            if uid is not None:
                await session.execute(
                    text("SELECT set_config('app.current_user_id', :uid, true)"), {"uid": str(uid)}
                )
            yield session

    settings = get_settings()

    def headers(uid):
        token = jwt.encode(
            {
                "sub": str(uid),
                "aud": settings.JWT_AUDIENCE,
                "iss": settings.JWT_ISSUER,
                "exp": int(time.time()) + 300,
            },
            pair,
            algorithm="RS256",
            headers={"kid": "site-filter-test"},
        )
        return {"Authorization": "Bearer " + token}

    previous = app.dependency_overrides.copy()
    app.dependency_overrides.update({get_db: db, get_db_snapshot: db})
    try:
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://api.test"
        ) as client:
            yield client, headers
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)


async def assert_aggregate_population(client, auth, params, expected):
    """Expected counts come from declared fixture sets, not another server response."""
    graph = await client.get("/api/v1/painel/filtrado/metricas", params=params, headers=auth)
    assert graph.status_code == 200
    assert graph.json()["total_periodo"]["total_apostas"] == expected
    assert sum(graph.json()["datasets"][2]["data"]) == expected
    export = await client.get("/api/v1/painel/filtrado/export", params=params, headers=auth)
    assert export.status_code == 200
    book = load_workbook(io.BytesIO(export.content), read_only=True)
    try:
        for name in book.sheetnames[:5]:
            rows = list(book[name].values)
            index = rows[0].index("total_apostas")
            assert sum(row[index] for row in rows[1:]) == expected
    finally:
        book.close()


@pytest.mark.parametrize(
    "visibilidade,indices", [("ativas", range(10)), ("apagadas", [10, 11]), ("todas", range(12))]
)
async def test_visibility_pagination_and_empty_page_match_independent_sets(
    api, dataset, visibilidade, indices
):
    client, headers = api
    expected = {dataset.keys[i] for i in indices}
    found = set()
    for page in range(1, 8):
        r = await client.get(
            "/api/v1/apostas",
            params={"visibilidade": visibilidade, "page_size": 2, "page": page},
            headers=headers(dataset.user),
        )
        assert r.status_code == 200
        assert r.json()["pagination"]["total"] == len(expected)
        found.update(a["chave"] for a in r.json()["data"])
    assert found == expected
    summary = await client.get(
        "/api/v1/painel/filtrado",
        params={"visibilidade": visibilidade},
        headers=headers(dataset.user),
    )
    assert summary.status_code == 200 and summary.json()["resumo"]["total_apostas"] == len(expected)
    await assert_aggregate_population(
        client, headers(dataset.user), {"visibilidade": visibilidade}, len(expected)
    )


@pytest.mark.parametrize(
    "campo,dimension,expected",
    [
        ("casa_id", "casas", [0, 2, 4, 6, 8]),
        ("titular_id", "titulares", [0, 2, 4, 6, 8]),
        ("conta_casa_id", "contas", [0, 2, 4, 6, 8]),
        ("tipster_id", "tipsters", [0, 2, 4, 6, 8]),
        ("mercado_id", "mercados", [0, 2, 4, 6, 8]),
        ("competicao_id", "competicoes", [0, 2, 4, 6, 8]),
        ("banca_id", "bancas", [0, 2, 4, 6, 8]),
        ("grupo_id", "grupos", list(range(10))),
    ],
)
async def test_each_dimension_selects_before_page_and_without_duplicate_group_facts(
    api, dataset, campo, dimension, expected
):
    client, headers = api
    id = getattr(dataset, dimension)[0]
    params = {campo: str(id), "page_size": 100}
    result = (
        await client.get("/api/v1/apostas", params=params, headers=headers(dataset.user))
    ).json()
    assert {v["chave"] for v in result["data"]} == {dataset.keys[i] for i in expected}
    assert result["pagination"]["total"] == len(expected)
    summary = (
        await client.get(
            "/api/v1/painel/filtrado", params={campo: str(id)}, headers=headers(dataset.user)
        )
    ).json()["resumo"]
    assert summary["total_apostas"] == len(expected)
    options = await client.get(
        "/api/v1/filtros/" + dimension, params={"id": str(id)}, headers=headers(dataset.user)
    )
    assert options.json()["data"][0]["id"] == str(id)
    await assert_aggregate_population(
        client, headers(dataset.user), {campo: str(id)}, len(expected)
    )


@pytest.mark.parametrize(
    "params,indices",
    [
        ({"estado": "GREEN"}, [0, 4, 8, 9]),
        ({"origem": "manual"}, [0, 2, 4, 6, 8]),
        ({"revisao_grave": "true"}, [8]),
        ({"estado": "GREEN", "origem": "manual", "revisao_grave": "false"}, [0, 4]),
        (
            {"desde": "2026-09-10T00:00:00-03:00", "ate": "2026-09-11T00:00:00-03:00"},
            list(range(1, 10)),
        ),
        ({"desde": "2026-09-10T03:00:00Z", "ate": "2026-09-10T03:00:01Z"}, list(range(1, 10))),
    ],
)
async def test_state_origin_review_and_timezone_boundaries(api, dataset, params, indices):
    client, headers = api
    r = await client.get(
        "/api/v1/apostas", params={**params, "page_size": 100}, headers=headers(dataset.user)
    )
    assert r.status_code == 200
    assert {v["chave"] for v in r.json()["data"]} == {dataset.keys[i] for i in indices}
    p = await client.get("/api/v1/painel/filtrado", params=params, headers=headers(dataset.user))
    assert p.status_code == 200 and p.json()["resumo"]["total_apostas"] == len(indices)
    await assert_aggregate_population(client, headers(dataset.user), params, len(indices))


@pytest.mark.parametrize(
    "campo,dim",
    [
        ("titular_id", "titulares"),
        ("conta_casa_id", "contas"),
        ("grupo_id", "grupos"),
        ("banca_id", "bancas"),
    ],
)
async def test_private_foreign_ids_and_nonexistent_ids_are_indistinguishable(
    api, dataset, campo, dim
):
    client, headers = api
    for id in [dataset.other_ids[campo], 9223372036854775807]:
        for path in [
            "/api/v1/apostas",
            "/api/v1/painel/filtrado",
            "/api/v1/painel/filtrado/metricas",
        ]:
            r = await client.get(path, params={campo: str(id)}, headers=headers(dataset.user))
            assert r.status_code == 200
            body = r.json()
            assert (
                body["pagination"]["total"]
                if path.endswith("apostas")
                else body.get("resumo", body.get("total_periodo"))["total_apostas"]
            ) == 0
        options = await client.get(
            "/api/v1/filtros/" + dim, params={"id": str(id)}, headers=headers(dataset.user)
        )
        assert (
            options.status_code == 200
            and options.json()["data"] == []
            and options.json()["pagination"]["total"] == 0
        )
        await assert_aggregate_population(client, headers(dataset.user), {campo: str(id)}, 0)


async def test_rls_without_user_and_cross_owner_group_fk(engine_app, como, dataset):
    async with como(engine_app, None) as session:
        assert await session.scalar(select(func.count()).select_from(models.GrupoAposta)) == 0
        assert await session.scalar(select(func.count()).select_from(models.ApostaGrupo)) == 0
    async with como(engine_app, dataset.user) as session:
        assert (
            await session.scalar(
                select(models.GrupoAposta.id).where(
                    models.GrupoAposta.id == dataset.other_ids["grupo_id"]
                )
            )
            is None
        )
        own = await session.scalar(
            select(models.Aposta.id).where(models.Aposta.chave == dataset.keys[0])
        )
        session.add(
            models.ApostaGrupo(
                usuario_id=dataset.user, aposta_id=own, grupo_id=dataset.other_ids["grupo_id"]
            )
        )
        from sqlalchemy.exc import IntegrityError

        with pytest.raises(IntegrityError):
            await session.flush()


async def test_financial_summary_daily_graphs_and_all_xlsx_sections_use_independent_amounts(
    api, dataset
):
    client, headers = api
    params = {
        "desde": "2026-09-10T02:59:59Z",
        "ate": "2026-09-10T03:00:01Z",
        "revisao_grave": "false",
        "estado": "MEIO_GREEN",
    }
    summary = (
        await client.get("/api/v1/painel/filtrado", params=params, headers=headers(dataset.user))
    ).json()["resumo"]
    assert {
        k: summary[k]
        for k in [
            "total_apostas",
            "giro_centavos",
            "base_roi_centavos",
            "retorno_centavos",
            "lucro_centavos",
        ]
    } == {
        "total_apostas": 1,
        "giro_centavos": 100,
        "base_roi_centavos": 100,
        "retorno_centavos": 150,
        "lucro_centavos": 50,
    }
    graph = (
        await client.get(
            "/api/v1/painel/filtrado/metricas", params=params, headers=headers(dataset.user)
        )
    ).json()
    assert graph["labels"] == ["2026-09-10"] and graph["granularidade"] == "dia"
    assert graph["datasets"][0]["data"] == [50]
    response = await client.get(
        "/api/v1/painel/filtrado/export", params=params, headers=headers(dataset.user)
    )
    assert response.status_code == 200 and response.headers["cache-control"] == "private, no-store"
    book = load_workbook(io.BytesIO(response.content), read_only=True)
    assert book.sheetnames == [
        "Resumo",
        "Por casa",
        "Por tipster",
        "Por mercado",
        "Por periodo",
        "Evolucao",
    ]
    for name in book.sheetnames[:5]:
        rows = list(book[name].values)
        values = dict(
            zip(
                rows[0],
                (*rows[1], *(None for _ in range(len(rows[0]) - len(rows[1])))),
                strict=True,
            )
        )
        assert values["total_apostas"] == 1 and values["lucro_centavos"] == 50
        assert values["giro_centavos"] == 100 and values["retorno_centavos"] == 150
    rows = list(book["Evolucao"].values)
    values = dict(
        zip(rows[0], (*rows[1], *(None for _ in range(len(rows[0]) - len(rows[1])))), strict=True)
    )
    assert values["contribuicao_centavos"] == values["acumulado_centavos"] == 50
    assert values["saldo_centavos"] is None
    book.close()


async def test_unknown_and_missing_dates_are_explicit_not_fabricated_zero(api, dataset):
    client, headers = api
    summary = (await client.get("/api/v1/painel/filtrado", headers=headers(dataset.user))).json()[
        "resumo"
    ]
    assert summary["total_apostas"] == 10 and summary["pendentes"] == 1 and summary["anuladas"] == 1
    assert summary["em_revisao"] == 1 and summary["resultados_desconhecidos"] == 1
    assert summary["giro_centavos"] == 600 and summary["base_roi_centavos"] == 700
    assert (
        summary["retorno_centavos"] is None
        and summary["lucro_centavos"] is None
        and summary["roi"] is None
    )
    graph = (
        await client.get(
            "/api/v1/painel/filtrado/metricas",
            params={"visibilidade": "todas"},
            headers=headers(dataset.user),
        )
    ).json()
    assert graph["labels"] == ["2026-09-09", "2026-09-10", None]
    assert graph["datasets"][0]["data"] == [100, None, -100]
    assert graph["datasets"][2]["data"] == [1, 9, 2]
    export = await client.get("/api/v1/painel/filtrado/export", headers=headers(dataset.user))
    assert export.status_code == 200
    book = load_workbook(io.BytesIO(export.content), read_only=True)
    try:
        headings = next(book["Resumo"].values)
        row = next(
            book["Resumo"].iter_rows(min_row=2, max_row=2, max_col=len(headings), values_only=True)
        )
        amounts = dict(zip(headings, row, strict=True))
        assert all(amounts[k] is None for k in ("retorno_centavos", "lucro_centavos", "roi"))
        assert amounts["giro_centavos"] == 600 and amounts["base_roi_centavos"] == 700
    finally:
        book.close()


async def test_catalog_page_search_selected_inactive_and_exact_bigint(
    api, dataset, engine_app, como
):
    client, headers = api
    for dim, id in [
        ("casas", dataset.casas[1]),
        ("contas", dataset.contas[1]),
        ("titulares", dataset.titulares[1]),
        ("grupos", dataset.grupos[1]),
    ]:
        response = await client.get(
            "/api/v1/filtros/" + dim, params={"id": str(id)}, headers=headers(dataset.user)
        )
        assert response.status_code == 200 and response.json()["data"][0]["ativa"] is False
        normal = (
            await client.get(
                "/api/v1/filtros/" + dim, params={"page_size": 1}, headers=headers(dataset.user)
            )
        ).json()
        assert str(id) not in {v["id"] for v in normal["data"]}
    id = 9007199254740993
    async with como(engine_app, dataset.user) as session:
        session.add(models.GrupoAposta(id=id, usuario_id=dataset.user, nome="Z id preciso"))
        await session.commit()
    response = await client.get(
        "/api/v1/filtros/grupos", params={"id": str(id)}, headers=headers(dataset.user)
    )
    assert response.json()["data"][0]["id"] == "9007199254740993"
    edited = await client.patch(
        "/api/v1/grupos/9007199254740993",
        json={"nome": "Z id preciso arquivado", "arquivado": True},
        headers=headers(dataset.user),
    )
    assert edited.status_code == 200
    assert edited.json() == {
        "id": "9007199254740993",
        "nome": "Z id preciso arquivado",
        "ativa": False,
    }
    assert (
        await client.patch(
            "/api/v1/grupos/9223372036854775808",
            json={"nome": "Inválido"},
            headers=headers(dataset.user),
        )
    ).status_code == 422
    missing_page = await client.get(
        "/api/v1/filtros/grupos",
        params={"page": 100, "page_size": 1},
        headers=headers(dataset.user),
    )
    # Archiving the exact BIGINT leaves only the fixture's active group in the catalog.
    assert missing_page.json()["data"] == [] and missing_page.json()["pagination"]["total"] == 1
    injected = await client.get(
        "/api/v1/filtros/grupos", params={"q": "' OR 1=1 --"}, headers=headers(dataset.user)
    )
    assert injected.json()["data"] == []
    exact = await client.get(
        "/api/v1/apostas", params={"grupo_id": str(id)}, headers=headers(dataset.user)
    )
    assert exact.status_code == 200 and exact.json()["pagination"]["total"] == 0


@pytest.mark.parametrize(
    "params",
    [
        {"visibilidade": "apagadas", "incluir_apagadas": "true"},
        {"visibilidade": "todas", "incluir_apagadas": "false"},
        {"visibilidade": "invalida"},
        {"grupo_id": "0"},
        {"banca_id": "9223372036854775808"},
        {"desde": "2026-09-11T00:00:00Z", "ate": "2026-09-10T00:00:00Z"},
        {"periodo": "all", "desde": "2020-01-01T00:00:00Z"},
    ],
)
async def test_invalid_or_conflicting_filters_are_not_silently_removed(api, dataset, params):
    client, headers = api
    for path in [
        "/api/v1/apostas",
        "/api/v1/painel/filtrado",
        "/api/v1/painel/filtrado/metricas",
        "/api/v1/painel/filtrado/export",
    ]:
        assert (
            await client.get(path, params=params, headers=headers(dataset.user))
        ).status_code == 422


async def test_legacy_queries_default_and_all_remain_compatible(api, dataset):
    client, headers = api
    for value, expected in [(None, 10), ("false", 10), ("true", 12)]:
        params = {} if value is None else {"incluir_apagadas": value}
        response = await client.get("/api/v1/apostas", params=params, headers=headers(dataset.user))
        assert response.status_code == 200 and response.json()["pagination"]["total"] == expected
    for path in [
        "/api/v1/filtros/grupos",
        "/api/v1/painel/filtrado",
        "/api/v1/painel/filtrado/metricas",
        "/api/v1/painel/filtrado/export",
    ]:
        assert (await client.get(path)).status_code == 401


async def test_delete_restore_preserve_identity_events_and_groups(api, dataset, engine_app, como):
    client, headers = api
    key = dataset.keys[0]
    async with como(engine_app, dataset.user) as session:
        bet = await session.scalar(select(models.Aposta).where(models.Aposta.chave == key))
        id = bet.id
        casa_nome = await session.scalar(
            select(models.Casa.nome).where(models.Casa.id == dataset.casas[0])
        )
        await EventoRepo().append(
            session,
            {
                "usuario_id": dataset.user,
                "tipo": "APOSTA_CRIADA",
                "fonte": "manual",
                "aposta_chave": key,
                "payload_json": {
                    "origem": "manual",
                    "casa": casa_nome,
                    "odd": 2,
                    "stake_unidades": 1,
                    "valor_unidade_centavos": 100,
                    "data_aposta": "2026-09-10T02:59:59+00:00",
                    "data_jogo": "2026-10-01T00:00:00+00:00",
                    "conta_casa_id": dataset.contas[0],
                    "conta_referencia_explicita": True,
                },
            },
        )
        await session.commit()
    assert (
        await client.delete("/api/v1/apostas/" + key, headers=headers(dataset.user))
    ).status_code == 200
    deleted = (
        await client.get(
            "/api/v1/apostas",
            params={"visibilidade": "apagadas", "page_size": 100},
            headers=headers(dataset.user),
        )
    ).json()
    assert key in {v["chave"] for v in deleted["data"]}
    assert (
        await client.post("/api/v1/apostas/" + key + "/restaurar", headers=headers(dataset.user))
    ).status_code == 200
    async with como(engine_app, dataset.user) as session:
        bet = await session.scalar(select(models.Aposta).where(models.Aposta.chave == key))
        assert bet.id == id and bet.selecionada is True
        assert (
            await session.scalar(
                select(func.count())
                .select_from(models.ApostaGrupo)
                .where(models.ApostaGrupo.aposta_id == id)
            )
            == 2
        )
        types = list(
            (
                await session.execute(
                    select(models.Evento.tipo).where(models.Evento.aposta_chave == key)
                )
            ).scalars()
        )
        assert "APOSTA_CRIADA" in types and len(types) >= 3


@pytest.mark.parametrize(
    "estado,origem,total,giro,base,retorno,lucro",
    [
        ("GREEN", "manual", 2, 100, 200, 500, 400),
        ("RED", None, 1, 100, 100, 0, -100),
        ("PENDENTE", None, 1, 0, 0, 0, 0),
        ("ANULADA", None, 1, 0, 0, 0, 0),
        ("MEIO_GREEN", None, 1, 100, 100, 150, 50),
        ("MEIO_RED", None, 1, 100, 100, 50, -50),
        ("CASHOUT", None, 1, 100, 100, 125, 25),
    ],
)
async def test_financial_cases_have_independent_expected_values(
    api, dataset, estado, origem, total, giro, base, retorno, lucro
):
    client, headers = api
    params = {"estado": estado, "revisao_grave": "false"}
    if origem is not None:
        params["origem"] = origem
    summary = (
        await client.get("/api/v1/painel/filtrado", params=params, headers=headers(dataset.user))
    ).json()["resumo"]
    assert [
        summary[k]
        for k in (
            "total_apostas",
            "giro_centavos",
            "base_roi_centavos",
            "retorno_centavos",
            "lucro_centavos",
        )
    ] == [total, giro, base, retorno, lucro]
    graph = (
        await client.get(
            "/api/v1/painel/filtrado/metricas", params=params, headers=headers(dataset.user)
        )
    ).json()
    assert sum(graph["datasets"][0]["data"]) == lucro
    assert sum(graph["datasets"][1]["data"]) == giro
    assert sum(graph["datasets"][2]["data"]) == total


async def test_group_commands_archive_historical_links_and_audit_atomically(
    api, dataset, engine_app, como
):
    client, headers = api
    auth = headers(dataset.user)
    created = await client.post("/api/v1/grupos", json={"nome": " Grupo criado "}, headers=auth)
    assert created.status_code == 201 and created.json()["nome"] == "Grupo criado"
    id = created.json()["id"]
    assert (
        await client.post("/api/v1/grupos", json={"nome": "Grupo criado"}, headers=auth)
    ).status_code == 422
    path = "/api/v1/apostas/" + dataset.keys[0] + "/grupos"
    for ids in (
        [id, id],
        [str(dataset.other_ids["grupo_id"])],
        [str(dataset.grupos[1])],
        ["9223372036854775808"],
    ):
        assert (await client.put(path, json={"grupo_ids": ids}, headers=auth)).status_code == 422
    assert (
        await client.patch(
            "/api/v1/grupos/" + str(dataset.other_ids["grupo_id"]),
            json={"nome": "Invasão"},
            headers=auth,
        )
    ).status_code == 404
    assert (await client.put(path, json={"grupo_ids": [id]}, headers=auth)).status_code == 200
    assert (
        await client.patch(
            "/api/v1/grupos/" + id,
            json={"nome": "Grupo arquivado", "arquivado": True},
            headers=auth,
        )
    ).status_code == 200
    options = (await client.get("/api/v1/filtros/grupos", params={"id": id}, headers=auth)).json()
    assert options["data"] == [{"id": id, "nome": "Grupo arquivado", "ativa": False}]
    selected = (await client.get("/api/v1/apostas", params={"grupo_id": id}, headers=auth)).json()
    assert {row["chave"] for row in selected["data"]} == {dataset.keys[0]}
    assert (await client.put(path, json={"grupo_ids": [id]}, headers=auth)).status_code == 422
    async with como(engine_app, dataset.user) as session:
        audits = list(
            (
                await session.execute(
                    select(models.AuditLog.action).where(
                        models.AuditLog.usuario_id == dataset.user,
                        models.AuditLog.resource_type == "grupos_aposta",
                        models.AuditLog.resource_id == id,
                    )
                )
            ).scalars()
        )
        assert audits.count("INSERT") == audits.count("UPDATE") == 1
        link = await session.scalar(
            select(models.AuditLog).where(
                models.AuditLog.resource_type == "apostas_grupos",
                models.AuditLog.action == "INSERT",
                models.AuditLog.diff["grupo_id"].astext == id,
            )
        )
        assert (
            link is not None
            and link.actor_usuario_id == dataset.user
            and link.resource_id.endswith(":" + id)
        )
    assert (await client.put(path, json={"grupo_ids": []}, headers=auth)).status_code == 200
    assert (await client.get("/api/v1/apostas", params={"grupo_id": id}, headers=auth)).json()[
        "pagination"
    ]["total"] == 0


async def test_gets_are_read_only_and_origins_are_actual_domain_ids(api, dataset, engine_app, como):
    client, headers = api

    async def facts():
        async with como(engine_app, dataset.user) as session:
            return [
                await session.scalar(select(func.count()).select_from(model))
                for model in (
                    models.Aposta,
                    models.Evento,
                    models.AuditLog,
                    models.ApostaGrupo,
                    models.GrupoAposta,
                )
            ]

    before = await facts()
    for dimension in (
        "casas",
        "tipsters",
        "mercados",
        "competicoes",
        "origens",
        "grupos",
        "bancas",
        "titulares",
        "contas",
    ):
        r = await client.get("/api/v1/filtros/" + dimension, headers=headers(dataset.user))
        assert r.status_code == 200 and r.headers["vary"] == "Authorization, Cookie"
        if dimension == "origens":
            assert {v["id"] for v in r.json()["data"]} == {
                "telegram",
                "telegram_bot",
                "print",
                "manual",
                "planilha",
                "casa",
            }
    await assert_aggregate_population(
        client, headers(dataset.user), {"grupo_id": str(dataset.grupos[0])}, 10
    )
    assert before == await facts()


async def test_explicit_bank_correction_is_owned_audited_and_replayable(
    api, dataset, engine_app, como
):
    from bancaemdia.domain.projecao import projetar

    client, headers = api
    key = dataset.keys[0]
    async with como(engine_app, dataset.user) as session:
        casa = await session.scalar(
            select(models.Casa.nome).where(models.Casa.id == dataset.casas[0])
        )
        await EventoRepo().append(
            session,
            {
                "usuario_id": dataset.user,
                "tipo": "APOSTA_CRIADA",
                "fonte": "manual",
                "aposta_chave": key,
                "payload_json": {
                    "origem": "manual",
                    "casa": casa,
                    "odd": 2,
                    "stake_unidades": 1,
                    "valor_unidade_centavos": 100,
                    "data_aposta": "2026-09-10T02:59:59Z",
                    "data_jogo": "2026-10-01T00:00:00Z",
                    "conta_casa_id": dataset.contas[0],
                    "conta_referencia_explicita": True,
                    "banca_id": dataset.bancas[0],
                },
            },
        )
        await session.commit()
    path = "/api/v1/apostas/" + key
    assert (
        await client.patch(
            path, json={"banca_id": dataset.other_ids["banca_id"]}, headers=headers(dataset.user)
        )
    ).status_code == 422
    assert (
        await client.patch(
            path, json={"banca_id": dataset.bancas[1]}, headers=headers(dataset.user)
        )
    ).status_code == 200
    async with como(engine_app, dataset.user) as session:
        bet = await session.scalar(select(models.Aposta).where(models.Aposta.chave == key))
        assert bet.banca_id == dataset.bancas[1] and bet.conta_casa_id == dataset.contas[0]
        events = list(
            (
                await session.execute(
                    select(models.Evento)
                    .where(models.Evento.aposta_chave == key)
                    .order_by(models.Evento.id)
                )
            ).scalars()
        )
        assert len(events) == 2
        folded, _ = projetar([(e.tipo, e.fonte, e.payload_json) for e in events])
        assert folded["banca_id"] == dataset.bancas[1]
    assert (
        await client.patch(path, json={"banca_id": None}, headers=headers(dataset.user))
    ).status_code == 200
    async with como(engine_app, dataset.user) as session:
        assert (
            await session.scalar(select(models.Aposta).where(models.Aposta.chave == key))
        ).banca_id is None


async def test_catalog_database_timeout_is_503_not_a_successful_empty_list(
    api, dataset, monkeypatch
):
    client, headers = api
    original = AsyncSession.execute

    async def timed_out(session, statement, *args, **kwargs):
        if "grupos_aposta" in str(statement):
            # Exercise a real PostgreSQL statement timeout, not a mocked result/error.
            await original(session, text("SET LOCAL statement_timeout = '1ms'"))
            return await original(session, text("SELECT pg_sleep(0.05)"))
        return await original(session, statement, *args, **kwargs)

    monkeypatch.setattr(AsyncSession, "execute", timed_out)
    response = await client.get("/api/v1/filtros/grupos", headers=headers(dataset.user))
    assert response.status_code == 503
    assert response.json() == {"detail": "catálogo temporariamente indisponível"}
    assert response.headers["cache-control"] == "private, no-store"
    assert response.headers["vary"] == "Authorization, Cookie"


@pytest.mark.parametrize(
    "house_selected,visibility,count,financial",
    [
        (True, "ativas", 1, 100),
        (True, "apagadas", 0, 0),
        (True, "todas", 2, 100),
        (False, "ativas", 0, 0),
        (False, "apagadas", 1, 100),
        (False, "todas", 2, 100),
    ],
)
async def test_consolidated_sources_preserve_visibility_without_double_financial_count(
    api, dataset, engine_app, como, house_selected, visibility, count, financial
):
    async with como(engine_app, dataset.user) as session:
        group = models.GrupoAposta(usuario_id=dataset.user, nome="Consolidação " + uuid4().hex)
        session.add(group)
        bets = []
        for origin, selected in [("casa", house_selected), ("telegram", True)]:
            bet = models.Aposta(
                usuario_id=dataset.user,
                chave=uuid4().hex,
                origem=origin,
                estado="GREEN",
                selecionada=selected,
                chat_id=123 if origin == "telegram" else None,
                message_id=456 if origin == "telegram" else None,
                conta_casa_id=dataset.contas[0],
                stake_centavos=100,
                stake_unidades=1,
                valor_aposta_centavos=100,
                retorno_centavos=200,
                data_aposta=datetime(2026, 9, 10, tzinfo=UTC),
            )
            session.add(bet)
            bets.append(bet)
        await session.flush()
        session.add(
            models.ApostaConsolidacao(
                usuario_id=dataset.user,
                casa_aposta_id=bets[0].id,
                telegram_aposta_id=bets[1].id,
                conta_casa_id=dataset.contas[0],
                estado="active",
                decisao="reviewed",
                versao="site-filter-acceptance",
                evidencia={},
                contexto={},
                ator="test",
            )
        )
        session.add_all([
            models.ApostaGrupo(usuario_id=dataset.user, aposta_id=bet.id, grupo_id=group.id)
            for bet in bets
        ])
        gid = group.id
        await session.commit()
    client, headers = api
    auth = headers(dataset.user)
    params = {"grupo_id": str(gid), "visibilidade": visibility}
    listing = await client.get("/api/v1/apostas", params=params, headers=auth)
    assert listing.status_code == 200
    assert listing.json()["pagination"]["total"] == count
    summary = await client.get("/api/v1/painel/filtrado", params=params, headers=auth)
    assert summary.status_code == 200
    expected = summary.json()["resumo"]
    assert expected["total_apostas"] == count
    assert expected["giro_centavos"] == expected["lucro_centavos"] == financial
    assert expected["retorno_centavos"] == financial * 2
    graph = await client.get("/api/v1/painel/filtrado/metricas", params=params, headers=auth)
    assert graph.status_code == 200
    assert graph.json()["total_periodo"]["giro_centavos"] == financial
    await assert_aggregate_population(client, auth, params, count)
    export = await client.get("/api/v1/painel/filtrado/export", params=params, headers=auth)
    assert export.status_code == 200
    book = load_workbook(io.BytesIO(export.content), read_only=True)
    try:
        for name in book.sheetnames[:5]:
            rows = list(book[name].values)
            stake_index = rows[0].index("giro_centavos")
            assert sum(row[stake_index] for row in rows[1:]) == financial
    finally:
        book.close()


async def test_read_only_access_keeps_site_gets_and_blocks_group_writes_in_database(
    api, dataset, engine_admin, engine_app, como
):
    # Activate only the disposable database guard; absence of a confirmed grant is
    # read-only. Restore the rollout timestamp without creating/changing subscriptions.
    async with engine_admin.begin() as conn:
        previous = await conn.scalar(text("SELECT activated_at FROM billing_rollout WHERE id=1"))
        await conn.execute(text("UPDATE billing_rollout SET activated_at=now() WHERE id=1"))
    try:
        client, headers = api
        auth = headers(dataset.user)
        for path in [
            "/api/v1/apostas",
            "/api/v1/painel/filtrado",
            "/api/v1/painel/filtrado/metricas",
            "/api/v1/painel/filtrado/export",
            "/api/v1/filtros/grupos",
            "/api/v1/filtros/bancas",
        ]:
            assert (await client.get(path, headers=auth)).status_code == 200
        for method, path, body in [
            ("POST", "/api/v1/grupos", {"nome": "Bloqueado"}),
            ("PATCH", "/api/v1/grupos/" + str(dataset.grupos[0]), {"nome": "Bloqueado"}),
            ("PUT", "/api/v1/apostas/" + dataset.keys[0] + "/grupos", {"grupo_ids": []}),
        ]:
            response = await client.request(method, path, json=body, headers=auth)
            assert response.status_code == 402
            assert response.json() == {"detail": "account_read_only"}
        for sql in [
            "INSERT INTO grupos_aposta(usuario_id,nome) VALUES (:uid,'Negado pelo banco')",
            "DELETE FROM apostas_grupos WHERE usuario_id=:uid",
        ]:
            async with como(engine_app, dataset.user) as session:
                with pytest.raises(DBAPIError) as error:
                    await session.execute(text(sql), {"uid": dataset.user})
                assert error.value.orig.sqlstate == "P0402"
                await session.rollback()
    finally:
        async with engine_admin.begin() as conn:
            await conn.execute(
                text("UPDATE billing_rollout SET activated_at=:previous WHERE id=1"),
                {"previous": previous},
            )
