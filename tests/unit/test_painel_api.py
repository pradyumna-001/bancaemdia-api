from __future__ import annotations

import asyncio
import io
from collections.abc import AsyncIterator, Mapping
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from bancaemdia.api.deps import get_current_user_snapshot
from bancaemdia.api.v1 import painel as rota
from bancaemdia.db.session import get_db_snapshot
from bancaemdia.domain.painel import (
    COLUNAS_EXPORTACAO,
    FiltrosPainel,
    FrescorPainel,
    GranularidadePainel,
    GrupoPainel,
    LinhaExportacao,
    MetricasGraficos,
    MetricasPainel,
    Painel,
    PeriodoPainel,
    PontoEvolucao,
    SaldoPainel,
    SecaoExportacao,
    SerieGrafico,
)
from bancaemdia.domain.registros import Usuario
from bancaemdia.exportacao.painel_xlsx import (
    gerar_painel_xlsx_assincrono as gerar_xlsx_real,
)
from bancaemdia.exportacao.painel_xlsx import remover_arquivo_temporario as remover_xlsx_real

AGORA = datetime(2026, 9, 21, 15, 0, tzinfo=UTC)
RESPONDIDO_EM = datetime(2026, 9, 21, 15, 0, 12, 345_600, tzinfo=UTC)


def _metricas(usuario_id: int) -> MetricasPainel:
    lucro = usuario_id * 100
    return MetricasPainel(
        total_apostas=3,
        pendentes=0,
        greens=2,
        reds=1,
        giro_centavos=3_000,
        base_roi_centavos=3_000,
        retorno_centavos=3_000 + lucro,
        lucro_centavos=lucro,
        freebets=0,
    )


def _frescor(*, replica: Decimal | None = Decimal("0.1254")) -> FrescorPainel:
    return FrescorPainel(
        atualizado_em=AGORA,
        idade_mv_segundos=Decimal("12.3456"),
        respondido_em=RESPONDIDO_EM,
        replica_atraso_segundos=replica,
    )


def _painel(usuario_id: int) -> Painel:
    metricas = _metricas(usuario_id)
    saldo = usuario_id * 100
    return Painel(
        resumo=metricas,
        saldo=SaldoPainel(saldo, saldo, 0),
        por_casa=(GrupoPainel(11, f"Casa {usuario_id}", metricas),),
        por_tipster=(GrupoPainel(12, f"Tipster {usuario_id}", metricas),),
        por_mercado=(GrupoPainel(13, f"Mercado {usuario_id}", metricas, "GOLS"),),
        evolucao=(
            PontoEvolucao(
                periodo_inicio=date(2026, 9, 20),
                contribuicao_centavos=usuario_id * 10,
                acumulado_centavos=usuario_id * 20,
                banca_id=21,
                banca_nome=f"Banca {usuario_id}",
                saldo_centavos=10_000 + usuario_id * 20,
            ),
        ),
        frescor=_frescor(),
    )


def _valor_exportado(coluna: str, usuario_id: int) -> object:
    if coluna == "periodo_inicio":
        return date(2026, 9, 20)
    if coluna == "granularidade":
        return "dia"
    if coluna == "id":
        return 11
    if coluna == "banca_id":
        return 21
    if coluna in {"nome", "banca_nome"}:
        return f"Usuario {usuario_id}"
    if coluna == "familia":
        return "GOLS"
    if coluna == "saldo_escopo":
        return "contas_casa_all_time"
    if coluna in {"roi", "win_rate"}:
        return Decimal("0.5")
    return usuario_id * 100


class RepositorioPainelFake:
    def __init__(self) -> None:
        self.consultas: list[tuple[int, FiltrosPainel]] = []
        self.metricas: list[tuple[int, FiltrosPainel]] = []
        self.frescores: list[int] = []
        self.exportacoes: list[tuple[int, FiltrosPainel, tuple[SecaoExportacao, ...]]] = []
        self.falhar_exportacao = False

    async def consultar(
        self,
        session: object,
        usuario_id: int,
        filtros: FiltrosPainel,
    ) -> Painel:
        del session
        self.consultas.append((usuario_id, filtros))
        return _painel(usuario_id)

    async def consultar_metricas(
        self,
        session: object,
        usuario_id: int,
        filtros: FiltrosPainel,
    ) -> MetricasGraficos:
        del session
        self.metricas.append((usuario_id, filtros))
        return MetricasGraficos(
            granularidade=filtros.janela.granularidade,
            labels=("2026-09-14", "2026-09-21"),
            datasets=(
                SerieGrafico("lucro_centavos", "Lucro", "centavos", (100, usuario_id * 100)),
                SerieGrafico("roi_basis_points", "ROI", "basis_points", (500, 2_500)),
            ),
        )

    async def frescor(
        self,
        session: object,
        usuario_id: int,
    ) -> FrescorPainel:
        del session
        self.frescores.append(usuario_id)
        return _frescor(replica=None)

    async def iterar_exportacao(
        self,
        session: object,
        usuario_id: int,
        filtros: FiltrosPainel,
        secoes: tuple[SecaoExportacao, ...],
    ) -> AsyncIterator[LinhaExportacao]:
        del session
        selecionadas = tuple(secoes)
        self.exportacoes.append((usuario_id, filtros, selecionadas))
        for secao in selecionadas:
            colunas = COLUNAS_EXPORTACAO[secao]
            yield LinhaExportacao(
                secao,
                {coluna: _valor_exportado(coluna, usuario_id) for coluna in colunas},
            )
            if self.falhar_exportacao:
                raise RuntimeError("cursor de exportacao falhou")


async def _sessao() -> AsyncIterator[object]:
    await asyncio.sleep(0)
    yield object()


def _cliente(monkeypatch, repositorio: RepositorioPainelFake, usuario_id: int) -> TestClient:
    def usuario() -> Usuario:
        return Usuario(
            usuario_id,
            f"u{usuario_id}@teste.local",
            f"Usuario {usuario_id}",
            AGORA,
            True,
        )

    monkeypatch.setattr(rota, "PainelRepo", lambda: repositorio)
    app = FastAPI()
    app.include_router(rota.router)
    app.dependency_overrides[get_current_user_snapshot] = usuario
    app.dependency_overrides[get_db_snapshot] = _sessao
    return TestClient(app)


def _contem_float(valor: object) -> bool:
    if isinstance(valor, float):
        return True
    if isinstance(valor, Mapping):
        return any(_contem_float(item) for item in valor.values())
    if isinstance(valor, (list, tuple)):
        return any(_contem_float(item) for item in valor)
    return False


def _assert_private(response) -> None:
    assert response.headers["cache-control"] == "private, no-store"
    assert response.headers["vary"] == "Authorization, Cookie"
    assert "age" not in response.headers


def test_main_endpoint_passes_filters_and_publishes_exact_numeric_contract(monkeypatch) -> None:
    repositorio = RepositorioPainelFake()
    cliente = _cliente(monkeypatch, repositorio, 7)

    resposta = cliente.get(
        "/api/v1/painel",
        params={
            "periodo": "90d",
            "casa_id": 11,
            "tipster_id": 12,
            "mercado_id": 13,
            "fresh": "true",
        },
    )

    assert resposta.status_code == 200
    _assert_private(resposta)
    assert len(repositorio.consultas) == 1
    usuario_id, filtros = repositorio.consultas[0]
    assert usuario_id == 7
    assert filtros.periodo is PeriodoPainel.NOVENTA_DIAS
    assert filtros.janela.granularidade is GranularidadePainel.SEMANA
    assert (filtros.casa_id, filtros.tipster_id, filtros.mercado_id) == (11, 12, 13)

    corpo = resposta.json()
    assert corpo["resumo"] == {
        "total_apostas": 3,
        "pendentes": 0,
        "greens": 2,
        "reds": 1,
        "giro_centavos": 3_000,
        "base_roi_centavos": 3_000,
        "retorno_centavos": 3_700,
        "lucro_centavos": 700,
        "freebets": 0,
        "roi": "0.233333",
        "win_rate": "0.666667",
        "roi_basis_points": 2_333,
        "win_rate_basis_points": 6_667,
    }
    assert corpo["saldo"] == {
        "saldo_total_centavos": 700,
        "saldo_conhecido_centavos": 700,
        "contas_saldo_desconhecido": 0,
        "escopo": "contas_casa_all_time",
    }
    assert corpo["por_casa"][0]["nome"] == "Casa 7"
    assert corpo["por_mercado"][0]["familia"] == "GOLS"
    assert corpo["evolucao"][0]["lucro_periodo_centavos"] == 70
    assert corpo["idade_mv_segundos"] == "12.346"
    assert corpo["replica_atraso_segundos"] == "0.125"
    assert corpo["replica_atraso_disponivel"] is True
    assert corpo["replica_atraso_estado"] == "disponivel"
    assert _contem_float(corpo) is False


def test_private_no_store_prevents_two_users_from_sharing_a_cached_financial_response(
    monkeypatch,
) -> None:
    repositorio = RepositorioPainelFake()
    cliente_ana = _cliente(monkeypatch, repositorio, 7)
    cliente_bia = _cliente(monkeypatch, repositorio, 9)

    ana_primeira = cliente_ana.get("/api/v1/painel")
    bia = cliente_bia.get("/api/v1/painel")
    ana_segunda = cliente_ana.get("/api/v1/painel")

    for resposta in (ana_primeira, bia, ana_segunda):
        assert resposta.status_code == 200
        _assert_private(resposta)
    assert ana_primeira.content == ana_segunda.content
    assert ana_primeira.content != bia.content
    assert ana_primeira.json()["resumo"]["lucro_centavos"] == 700
    assert bia.json()["resumo"]["lucro_centavos"] == 900
    assert [usuario_id for usuario_id, _ in repositorio.consultas] == [7, 9, 7]


def test_metrics_endpoint_passes_filters_and_returns_chart_contract_with_freshness(
    monkeypatch,
) -> None:
    repositorio = RepositorioPainelFake()
    cliente = _cliente(monkeypatch, repositorio, 7)

    resposta = cliente.get(
        "/api/v1/painel/metricas",
        params={"periodo": "1y", "casa_id": 11, "tipster_id": 12, "mercado_id": 13},
    )

    assert resposta.status_code == 200
    _assert_private(resposta)
    usuario_id, filtros = repositorio.metricas[0]
    assert usuario_id == 7
    assert filtros.periodo is PeriodoPainel.UM_ANO
    assert filtros.janela.granularidade is GranularidadePainel.MES
    assert (filtros.casa_id, filtros.tipster_id, filtros.mercado_id) == (11, 12, 13)
    assert repositorio.frescores == [7]
    assert resposta.json() == {
        "granularidade": "mes",
        "labels": ["2026-09-14", "2026-09-21"],
        "datasets": [
            {
                "chave": "lucro_centavos",
                "label": "Lucro",
                "unidade": "centavos",
                "data": [100, 700],
            },
            {
                "chave": "roi_basis_points",
                "label": "ROI",
                "unidade": "basis_points",
                "data": [500, 2_500],
            },
        ],
        "atualizado_em": "2026-09-21T15:00:00Z",
        "idade_mv_segundos": "12.346",
        "respondido_em": "2026-09-21T15:00:12.345600Z",
        "replica_atraso_segundos": None,
        "replica_atraso_disponivel": False,
        "replica_atraso_estado": "primario_ou_sem_telemetria",
    }


def test_export_passes_filters_returns_a_valid_xlsx_and_cleans_the_temporary_file(
    monkeypatch,
    tmp_path: Path,
) -> None:
    repositorio = RepositorioPainelFake()
    removidos: list[Path] = []

    async def gerar(abas, *, nome_download: str):
        return await gerar_xlsx_real(
            abas,
            nome_download=nome_download,
            diretorio_temporario=tmp_path,
        )

    def remover(caminho: str | Path) -> None:
        removidos.append(Path(caminho))
        remover_xlsx_real(caminho)

    monkeypatch.setattr(rota, "gerar_painel_xlsx_assincrono", gerar)
    monkeypatch.setattr(rota, "remover_arquivo_temporario", remover)
    cliente = _cliente(monkeypatch, repositorio, 7)

    resposta = cliente.get(
        "/api/v1/painel/export",
        params={"periodo": "all", "casa_id": 11, "tipster_id": 12, "mercado_id": 13},
    )

    assert resposta.status_code == 200
    _assert_private(resposta)
    assert resposta.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "painel-" in resposta.headers["content-disposition"]
    assert resposta.headers["x-content-type-options"] == "nosniff"
    assert resposta.content.startswith(b"PK")

    assert len(repositorio.exportacoes) == len(SecaoExportacao)
    assert [chamada[2] for chamada in repositorio.exportacoes] == [
        (secao,) for secao in SecaoExportacao
    ]
    for usuario_id, filtros, _ in repositorio.exportacoes:
        assert usuario_id == 7
        assert filtros.periodo is PeriodoPainel.TODO
        assert (filtros.casa_id, filtros.tipster_id, filtros.mercado_id) == (11, 12, 13)

    workbook = load_workbook(io.BytesIO(resposta.content), read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == list(rota.NOMES_DAS_ABAS.values())
        for secao, nome in rota.NOMES_DAS_ABAS.items():
            linhas = list(workbook[nome].iter_rows(values_only=True))
            assert linhas[0] == COLUNAS_EXPORTACAO[secao]
            assert len(linhas) == 2
    finally:
        workbook.close()

    assert len(removidos) == 1
    assert not removidos[0].exists()
    assert list(tmp_path.iterdir()) == []


def test_export_removes_a_partial_workbook_when_the_repository_stream_fails(
    monkeypatch,
    tmp_path: Path,
) -> None:
    repositorio = RepositorioPainelFake()
    repositorio.falhar_exportacao = True

    async def gerar(abas, *, nome_download: str):
        return await gerar_xlsx_real(
            abas,
            nome_download=nome_download,
            diretorio_temporario=tmp_path,
        )

    monkeypatch.setattr(rota, "gerar_painel_xlsx_assincrono", gerar)
    cliente = _cliente(monkeypatch, repositorio, 7)

    with pytest.raises(RuntimeError, match="cursor de exportacao falhou"):
        cliente.get("/api/v1/painel/export")

    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize(
    ("query", "campo"),
    [
        ({"periodo": "ontem"}, "periodo"),
        ({"casa_id": 0}, "casa_id"),
        ({"tipster_id": -1}, "tipster_id"),
        ({"mercado_id": 2**63}, "mercado_id"),
    ],
)
def test_invalid_filters_are_rejected_before_the_repository(
    monkeypatch,
    query: dict[str, object],
    campo: str,
) -> None:
    repositorio = RepositorioPainelFake()
    cliente = _cliente(monkeypatch, repositorio, 7)

    resposta = cliente.get("/api/v1/painel", params=query)

    assert resposta.status_code == 422
    assert any(erro["loc"][-1] == campo for erro in resposta.json()["detail"])
    assert repositorio.consultas == []
