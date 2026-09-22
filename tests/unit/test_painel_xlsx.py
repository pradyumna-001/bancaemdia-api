from __future__ import annotations

import asyncio
from pathlib import Path

import pytest
from openpyxl import load_workbook

from bancaemdia.exportacao import painel_xlsx
from bancaemdia.exportacao.painel_xlsx import (
    MIME_XLSX,
    AbaXlsx,
    AbaXlsxAssincrona,
    gerar_painel_xlsx,
    gerar_painel_xlsx_assincrono,
    remover_arquivo_temporario,
)


def test_xlsx_has_multiple_sheets_and_keeps_cents_as_integers(tmp_path: Path) -> None:
    arquivo = gerar_painel_xlsx(
        [
            AbaXlsx(
                "Resumo",
                ("metrica", "valor_centavos"),
                iter((("lucro", 12_345), ("saldo", -500))),
            ),
            AbaXlsx(
                "Por casa",
                ("casa", "lucro_centavos"),
                iter((("Betano", 2_000),)),
            ),
        ],
        diretorio_temporario=tmp_path,
    )

    workbook = load_workbook(arquivo.caminho, read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == ["Resumo", "Por casa"]
        assert list(workbook["Resumo"].iter_rows(values_only=True)) == [
            ("metrica", "valor_centavos"),
            ("lucro", 12_345),
            ("saldo", -500),
        ]
        centavos = workbook["Resumo"]["B2"]
        assert centavos.value == 12_345
        assert centavos.data_type == "n"
    finally:
        workbook.close()
        remover_arquivo_temporario(arquivo.caminho)

    assert arquivo.nome_download == "painel.xlsx"
    assert arquivo.media_type == MIME_XLSX
    assert arquivo.tamanho_bytes > 0


def test_formula_like_text_is_written_as_text_not_as_a_formula(tmp_path: Path) -> None:
    perigosos = ("=2+2", "+SUM(A1:A2)", "-10+20", "@SUM(A1:A2)")
    arquivo = gerar_painel_xlsx(
        [AbaXlsx("Dados", ("texto",), ((valor,) for valor in perigosos))],
        diretorio_temporario=tmp_path,
    )

    workbook = load_workbook(arquivo.caminho, read_only=True, data_only=False)
    try:
        celulas = [linha[0] for linha in workbook["Dados"].iter_rows()][1:]
        assert [celula.value for celula in celulas] == [f"'{valor}" for valor in perigosos]
        assert {celula.data_type for celula in celulas} == {"s"}
    finally:
        workbook.close()
        remover_arquivo_temporario(arquivo.caminho)


def test_rows_are_appended_before_the_iterator_advances(monkeypatch, tmp_path: Path) -> None:
    escritas: list[tuple[object, ...]] = []

    class Planilha:
        def append(self, linha) -> None:
            escritas.append(tuple(linha))

    class Livro:
        def create_sheet(self, title: str) -> Planilha:
            assert title == "Dados"
            return Planilha()

        def save(self, caminho: Path) -> None:
            caminho.write_bytes(b"xlsx")

        def close(self) -> None:
            return None

    def linhas():
        assert escritas == [("id",)]
        yield (1,)
        assert escritas[-1] == (1,)
        yield (2,)

    monkeypatch.setattr(painel_xlsx, "Workbook", lambda *, write_only: Livro())

    arquivo = gerar_painel_xlsx(
        [AbaXlsx("Dados", ("id",), linhas())],
        diretorio_temporario=tmp_path,
    )

    assert escritas == [("id",), (1,), (2,)]
    remover_arquivo_temporario(arquivo.caminho)


def test_an_empty_export_is_still_a_valid_workbook(tmp_path: Path) -> None:
    arquivo = gerar_painel_xlsx([], diretorio_temporario=tmp_path)

    workbook = load_workbook(arquivo.caminho, read_only=True)
    try:
        assert workbook.sheetnames == ["Painel"]
        assert list(workbook["Painel"].iter_rows(values_only=True)) == []
    finally:
        workbook.close()
        remover_arquivo_temporario(arquivo.caminho)


def test_cleanup_is_idempotent(tmp_path: Path) -> None:
    arquivo = gerar_painel_xlsx(
        [AbaXlsx("Resumo", ("valor",), ())],
        diretorio_temporario=tmp_path,
    )

    assert arquivo.caminho.exists()
    remover_arquivo_temporario(arquivo.caminho)
    remover_arquivo_temporario(arquivo.caminho)

    assert not arquivo.caminho.exists()


async def test_async_rows_are_consumed_incrementally_into_a_valid_xlsx(tmp_path: Path) -> None:
    avancos: list[int] = []

    async def linhas():
        for numero in range(3):
            await asyncio.sleep(0)
            avancos.append(numero)
            yield (numero, numero * 100)

    arquivo = await gerar_painel_xlsx_assincrono(
        [AbaXlsxAssincrona("Periodo", ("dia", "lucro_centavos"), linhas())],
        diretorio_temporario=tmp_path,
    )

    assert avancos == [0, 1, 2]
    workbook = load_workbook(arquivo.caminho, read_only=True)
    try:
        assert list(workbook["Periodo"].iter_rows(values_only=True)) == [
            ("dia", "lucro_centavos"),
            (0, 0),
            (1, 100),
            (2, 200),
        ]
    finally:
        workbook.close()
        remover_arquivo_temporario(arquivo.caminho)


async def test_async_export_removes_partial_file_when_the_stream_fails(tmp_path: Path) -> None:
    async def linhas():
        await asyncio.sleep(0)
        yield (1,)
        raise RuntimeError("banco caiu")

    with pytest.raises(RuntimeError, match="banco caiu"):
        await gerar_painel_xlsx_assincrono(
            [AbaXlsxAssincrona("Dados", ("id",), linhas())],
            diretorio_temporario=tmp_path,
        )

    assert list(tmp_path.iterdir()) == []
