from __future__ import annotations

import asyncio
from collections.abc import AsyncIterable, Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PREFIXOS_DE_FORMULA = ("=", "+", "-", "@")


@dataclass(frozen=True, slots=True)
class AbaXlsx:
    nome: str
    cabecalhos: Sequence[object]
    linhas: Iterable[Iterable[object]]


@dataclass(frozen=True, slots=True)
class AbaXlsxAssincrona:
    nome: str
    cabecalhos: Sequence[object]
    linhas: AsyncIterable[Iterable[object]]


@dataclass(frozen=True, slots=True)
class ArquivoXlsx:
    caminho: Path
    nome_download: str
    media_type: str
    tamanho_bytes: int


def _valor_seguro(valor: object) -> object:
    if isinstance(valor, str) and valor.startswith(PREFIXOS_DE_FORMULA):
        return f"'{valor}"
    return valor


def _linha_segura(linha: Iterable[object]) -> tuple[object, ...]:
    # Somente a linha corrente fica em memoria; o iteravel completo nunca e materializado.
    return tuple(_valor_seguro(valor) for valor in linha)


def remover_arquivo_temporario(caminho: str | Path) -> None:
    """Remove o XLSX depois que o FileResponse terminar de envia-lo."""
    Path(caminho).unlink(missing_ok=True)


def _arquivo_temporario(diretorio: str | Path | None) -> Path:
    with NamedTemporaryFile(
        prefix="bancaemdia-painel-",
        suffix=".xlsx",
        dir=diretorio,
        delete=False,
    ) as temporario:
        return Path(temporario.name)


def _resultado(caminho: Path, nome_download: str) -> ArquivoXlsx:
    return ArquivoXlsx(
        caminho=caminho,
        nome_download=nome_download,
        media_type=MIME_XLSX,
        tamanho_bytes=caminho.stat().st_size,
    )


def _abortar_workbook(workbook: Workbook) -> None:
    # No modo write_only cada aba mantem um gerador XML aberto ate save(). Se a fonte assíncrona
    # falhar antes disso, fechar explicitamente as abas evita temporários e geradores pendurados.
    for planilha in workbook.worksheets:
        if not planilha.closed:
            planilha.close()


def gerar_painel_xlsx(
    abas: Iterable[AbaXlsx],
    *,
    nome_download: str = "painel.xlsx",
    diretorio_temporario: str | Path | None = None,
) -> ArquivoXlsx:
    """Grava as abas sequencialmente em disco, mantendo apenas uma linha em memoria."""
    caminho = _arquivo_temporario(diretorio_temporario)

    workbook = Workbook(write_only=True)
    criou_aba = False
    try:
        for aba in abas:
            planilha = workbook.create_sheet(title=aba.nome)
            criou_aba = True
            if aba.cabecalhos:
                planilha.append(_linha_segura(aba.cabecalhos))
            for linha in aba.linhas:
                planilha.append(_linha_segura(linha))

        # Um XLSX precisa ter ao menos uma planilha visivel, inclusive para um painel vazio.
        if not criou_aba:
            workbook.create_sheet(title="Painel")

        workbook.save(caminho)
    except BaseException:
        _abortar_workbook(workbook)
        remover_arquivo_temporario(caminho)
        raise
    finally:
        workbook.close()

    return _resultado(caminho, nome_download)


async def gerar_painel_xlsx_assincrono(
    abas: Iterable[AbaXlsxAssincrona],
    *,
    nome_download: str = "painel.xlsx",
    diretorio_temporario: str | Path | None = None,
) -> ArquivoXlsx:
    """Consome streams do banco linha a linha e grava um XLSX temporario.

    A compactacao final e bloqueante e roda fora do event loop. Ate esse ponto, apenas a linha
    atual e as estruturas internas do modo ``write_only`` ficam em memoria.
    """

    caminho = _arquivo_temporario(diretorio_temporario)
    workbook = Workbook(write_only=True)
    criou_aba = False
    try:
        for aba in abas:
            planilha = workbook.create_sheet(title=aba.nome)
            criou_aba = True
            if aba.cabecalhos:
                planilha.append(_linha_segura(aba.cabecalhos))
            async for linha in aba.linhas:
                planilha.append(_linha_segura(linha))
        if not criou_aba:
            workbook.create_sheet(title="Painel")
        await asyncio.to_thread(workbook.save, caminho)
    except BaseException:
        _abortar_workbook(workbook)
        remover_arquivo_temporario(caminho)
        raise
    finally:
        workbook.close()
    return _resultado(caminho, nome_download)
