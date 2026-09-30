"""Bounded official-source fetches and fail-closed table ingestion."""

import csv
import io
import re
import unicodedata
from html.parser import HTMLParser
from urllib.parse import urljoin, urlsplit

import httpx
from openpyxl import load_workbook

from bancaemdia.coleta.catalogo import Observation, Regulatory, safe_url

SPA_BASE = "https://www.gov.br/fazenda/pt-br/composicao/orgaos/secretaria-de-premios-e-apostas/transparencia-ativa-processos-de-autorizacao-de-apostas-de-quota-fixa/"
NATIONAL = SPA_BASE + "empresas-autorizadas"
JUDICIAL = SPA_BASE + "autorizadas-por-determinacao-judicial"
MAX_SNAPSHOT = 10 * 1024 * 1024


class OfficialTables(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.rows: list[list[str]] = []
        self.links: list[str] = []
        self.link_labels: list[tuple[str, str]] = []
        self.anchor: tuple[str, list[str]] | None = None
        self.row: list[str] | None = None
        self.cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "a":
            href = dict(attrs).get("href")
            if href:
                self.links.append(href)
                self.anchor = (href, [])
        if tag == "tr":
            self.row = []
        elif tag in ("td", "th") and self.row is not None:
            self.cell = []
        elif tag in ("br", "li", "p") and self.cell is not None:
            self.cell.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag == "a" and self.anchor is not None:
            self.link_labels.append((self.anchor[0], "".join(self.anchor[1])))
            self.anchor = None
        if tag in ("td", "th") and self.cell is not None and self.row is not None:
            self.row.append("".join(self.cell))
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.rows.append(self.row)
            self.row = None
        elif tag in ("li", "p") and self.cell is not None:
            self.cell.append("\n")

    def handle_data(self, data: str) -> None:
        if self.anchor is not None:
            self.anchor[1].append(data)
        if self.cell is not None:
            self.cell.append(data)


def folded(value: str) -> str:
    return " ".join(
        ""
        .join(c for c in unicodedata.normalize("NFKD", value) if not unicodedata.combining(c))
        .lower()
        .split()
    )


def _parts(value: str) -> list[str]:
    return [p.strip().strip("• ") for p in re.split(r"[\n\r;]+", value) if p.strip().strip("• ")]


def parse_rows(
    rows: list[list[str]], situation: Regulatory, *, unresolved: list[dict[str, str]] | None = None
) -> list[Observation]:
    header: tuple[int, int, int] | None = None
    company = ""
    observations: list[Observation] = []
    for row in rows:
        names = [folded(c) for c in row]
        brand_col = next((i for i, v in enumerate(names) if v in ("marca", "marcas")), None)
        host_col = next(
            (i for i, v in enumerate(names) if v in ("dominio", "dominios", "hostname")), None
        )
        company_col = next(
            (
                i
                for i, v in enumerate(names)
                if v in ("empresa", "denominacao social da empresa", "razao social")
            ),
            None,
        )
        if brand_col is not None and host_col is not None and company_col is not None:
            header = company_col, brand_col, host_col
            continue
        if header is None or len(row) <= max(header):
            continue
        ci, bi, hi = header
        if row[ci].strip():
            company = " ".join(row[ci].split())
        brands, hosts = _parts(row[bi]), _parts(row[hi])
        if not brands and not hosts:
            continue
        if len(brands) != len(hosts) or not company:
            raise ValueError("official row has ambiguous brand/host mapping; review source format")
        for brand, host in zip(brands, hosts, strict=True):
            if folded(host) in ("a definir", "a ser definido", "nao informado", "-"):
                if unresolved is None:
                    raise ValueError(
                        "unresolved official domain must be preserved in the source report"
                    )
                unresolved.append({
                    "brand": brand,
                    "company": company,
                    "domain_text": host,
                    "reason": "official_domain_unresolved",
                })
                continue
            observations.append(
                Observation(brand=brand, hostname=host, company=company, situation=situation)
            )
    if header is None or not observations:
        raise ValueError("official authorization table unavailable or empty; refuse mass removal")
    keys = [o.key for o in observations]
    if len(keys) != len(set(keys)):
        raise ValueError("duplicate exact host mapping in official source")
    return observations


def parse_html(
    raw: bytes, situation: Regulatory, *, unresolved: list[dict[str, str]] | None = None
) -> list[Observation]:
    parser = OfficialTables()
    parser.feed(raw.decode("utf-8-sig"))
    return parse_rows(parser.rows, situation, unresolved=unresolved)


def parse_spreadsheet(
    raw: bytes, situation: Regulatory, *, unresolved: list[dict[str, str]] | None = None
) -> list[Observation]:
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise ValueError("missing authorization worksheet")
        rows = [[str(v or "") for v in row] for row in sheet.iter_rows(values_only=True)]
        return parse_rows(rows, situation, unresolved=unresolved)
    finally:
        workbook.close()


def parse_csv(
    raw: bytes, situation: Regulatory, *, unresolved: list[dict[str, str]] | None = None
) -> list[Observation]:
    text = raw.decode("utf-8-sig")
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    return parse_rows(
        list(csv.reader(io.StringIO(text), dialect)), situation, unresolved=unresolved
    )


def spreadsheet_link(raw: bytes, page_url: str) -> str:
    parser = OfficialTables()
    parser.feed(raw.decode("utf-8-sig"))
    # Prefer the actual Excel link; preserve both page and download snapshots separately.
    preferred = [
        link for link, label in parser.link_labels if "baixar arquivo excel" in folded(label)
    ]
    candidates = preferred or parser.links
    links = [
        urljoin(page_url, link)
        for link in candidates
        if urlsplit(link).path.lower().endswith(".xlsx")
    ]
    if not links:
        raise ValueError("official downloadable spreadsheet is absent")
    for link in links:
        if urlsplit(link).hostname == "www.gov.br":
            return safe_url(link)
    raise ValueError("spreadsheet origin is not gov.br")


def fetch_official(
    url: str, allowed_hosts: set[str], client: httpx.Client
) -> tuple[bytes, list[str]]:
    chain = []
    for _ in range(6):
        safe_url(url)
        if urlsplit(url).hostname not in allowed_hosts:
            raise ValueError("unreviewed official-source origin")
        if url in chain:
            raise ValueError("source redirect loop")
        chain.append(url)
        with client.stream("GET", url, follow_redirects=False) as response:
            if response.is_redirect:
                url = urljoin(url, response.headers["location"])
                continue
            response.raise_for_status()
            data = bytearray()
            for chunk in response.iter_bytes():
                data.extend(chunk)
                if len(data) > MAX_SNAPSHOT:
                    raise ValueError("official snapshot exceeds 10 MiB")
            return bytes(data), chain
    raise ValueError("too many official redirects")
