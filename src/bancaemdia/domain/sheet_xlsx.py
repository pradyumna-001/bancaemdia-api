"""Bounded XLSX inspection. Never evaluate formulas, macros or external links."""

from __future__ import annotations

import hashlib
import json
import math
import re
import time
from datetime import date, datetime
from io import BytesIO
from typing import Any
from zipfile import ZipFile

from defusedxml.ElementTree import iterparse
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

from bancaemdia.core.body_limits import PLANILHA_MAX_BYTES
from bancaemdia.domain.materializar import casa_canonica
from bancaemdia.domain.planilha_import import (
    OPTIONAL,
    REQUIRED,
    PlanilhaInvalidaError,
    _account,
    _booleano,
    _data,
    _numero,
    _texto,
)

MAX_SHEETS = 10
MAX_COLUMNS = 64
MAX_ROWS = 1000
MAX_CELLS = 64000
MAX_INFLATED = 20 * 1024 * 1024
MAX_SECONDS = 10
FIELDS = ("linha_id", *REQUIRED, *OPTIONAL)
REQUIRED_FIELDS = ("linha_id", *REQUIRED)
type Cell = str | int | float | bool | None


class SheetError(ValueError):
    def __init__(self, code: str, status: int = 422) -> None:
        self.code = code
        self.status = status
        super().__init__(code)


def digest(value: object) -> str:
    return hashlib.sha256(
        json.dumps(
            value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str
        ).encode()
    ).hexdigest()


def inspect_xlsx(content: bytes) -> dict[str, Any]:
    if not content:
        raise SheetError("empty_xlsx")
    if len(content) > PLANILHA_MAX_BYTES:
        raise SheetError("xlsx_bytes_limit", 413)
    deadline = time.monotonic() + MAX_SECONDS
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            names = [entry.filename for entry in entries]
            if len(entries) > 2000 or len(names) != len(set(names)):
                raise SheetError("xlsx_structure_limit", 413)
            if any(
                name.startswith(("/", "\\"))
                or ".." in name.split("/")
                or "vba" in name.lower()
                or "externallinks" in name.lower()
                for name in names
            ):
                raise SheetError("unsafe_xlsx")
            if not {"[Content_Types].xml", "xl/workbook.xml"} <= set(names):
                raise SheetError("invalid_xlsx")
            if sum(entry.file_size for entry in entries) > MAX_INFLATED or any(
                entry.file_size > max(entry.compress_size, 1) * 100 for entry in entries
            ):
                raise SheetError("xlsx_expansion_limit", 413)
            cells = rows = sheets = 0
            for entry in entries:
                if entry.flag_bits & 1:
                    raise SheetError("unsafe_xlsx")
                if not entry.filename.endswith((".xml", ".rels")):
                    continue
                data = archive.read(entry)
                upper = data.upper()
                if (
                    b"<!DOCTYPE" in upper
                    or b"<!ENTITY" in upper
                    or b'TARGETMODE="EXTERNAL"' in upper
                    or b"MACROENABLED" in upper
                ):
                    raise SheetError("unsafe_xlsx")
                worksheet = entry.filename.startswith("xl/worksheets/") and entry.filename.endswith(
                    ".xml"
                )
                if worksheet:
                    sheets += 1
                for _, element in iterparse(
                    BytesIO(data),
                    events=("end",),
                    forbid_dtd=True,
                    forbid_entities=True,
                    forbid_external=True,
                ):
                    tag = element.tag.rsplit("}", 1)[-1]
                    if (
                        element.attrib.get("TargetMode", "").lower() == "external"
                        or "macroenabled" in element.attrib.get("ContentType", "").lower()
                    ):
                        raise SheetError("unsafe_xlsx")
                    if worksheet and tag == "row":
                        rows += 1
                        if int(element.attrib.get("r", "0")) > MAX_ROWS + 1:
                            raise SheetError("xlsx_rows_limit", 413)
                    if worksheet and tag == "c":
                        cells += 1
                        match = re.fullmatch(r"([A-Z]+)[1-9][0-9]*", element.attrib.get("r", ""))
                        if match is None or column_index_from_string(match[1]) > MAX_COLUMNS:
                            raise SheetError("xlsx_columns_limit", 413)
                    element.clear()
                    if time.monotonic() > deadline:
                        raise SheetError("xlsx_time_limit", 413)
                    if cells > MAX_CELLS or rows > MAX_ROWS + MAX_SHEETS or sheets > MAX_SHEETS:
                        raise SheetError("xlsx_structure_limit", 413)
                if cells > MAX_CELLS or rows > MAX_ROWS + MAX_SHEETS or sheets > MAX_SHEETS:
                    raise SheetError("xlsx_structure_limit", 413)
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        result: dict[str, Any] = {"sheets": []}
        try:
            count = 0
            for sheet in book.worksheets:
                sheet.reset_dimensions()
                material: list[dict[str, Any]] = []
                headers: list[str | None] = []
                for number, row in enumerate(sheet.iter_rows(max_col=MAX_COLUMNS), start=1):
                    values: list[Cell] = []
                    formulas: list[int] = []
                    for column, cell in enumerate(row, start=1):
                        value = cell.value
                        if cell.data_type in {"f", "e"}:
                            formulas.append(column)
                        if isinstance(value, (datetime, date)):
                            value = value.isoformat()
                        if value is not None and not isinstance(value, (str, int, float, bool)):
                            value = str(value)
                        if isinstance(value, float) and not math.isfinite(value):
                            value = str(value)
                        if isinstance(value, str) and len(value) > 1000:
                            raise SheetError("xlsx_cell_limit", 413)
                        values.append(value)
                    while values and values[-1] is None:
                        values.pop()
                    if number == 1:
                        if formulas:
                            raise SheetError("unsafe_header")
                        headers = [str(v).strip() if v is not None else None for v in values]
                    elif values:
                        count += 1
                        material.append({"number": number, "values": values, "formulas": formulas})
                    if count > MAX_ROWS or time.monotonic() > deadline:
                        raise SheetError("xlsx_structure_limit", 413)
                result["sheets"].append({"name": sheet.title, "headers": headers, "rows": material})
        finally:
            book.close()
        if not count:
            raise SheetError("empty_xlsx")
        return result
    except SheetError:
        raise
    except Exception as error:
        raise SheetError("invalid_xlsx") from error


def default_mapping(sheet: dict[str, Any]) -> dict[str, int]:
    headers = [str(value or "").lower().strip() for value in sheet["headers"]]
    return {field: headers.index(field) + 1 for field in FIELDS if headers.count(field) == 1}


def validate_mapping(mapping: dict[str, int], width: int) -> list[dict[str, Any]]:
    errors = []
    for field in REQUIRED_FIELDS:
        if field not in mapping:
            errors.append({"field": field, "code": "missing_mapping"})
    for field, column in mapping.items():
        if field not in FIELDS or not 1 <= column <= width:
            errors.append({"field": field, "code": "invalid_mapping"})
        elif list(mapping.values()).count(column) > 1:
            errors.append({"field": field, "code": "duplicate_mapping"})
    return errors


def normalized_row(
    row: dict[str, Any], mapping: dict[str, int]
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    values = row["values"]
    data = {
        field: values[column - 1] if column <= len(values) else None
        for field, column in mapping.items()
    }
    errors: list[dict[str, Any]] = []
    normalized: dict[str, Any] = {}
    for field in mapping:
        value = data[field]
        if mapping[field] in row["formulas"]:
            errors.append({"field": field, "code": "formula_or_cell_error"})
            continue
        try:
            if field == "linha_id":
                if not isinstance(value, str) or not re.fullmatch(
                    r"[A-Za-z0-9][A-Za-z0-9._:-]{0,127}", value
                ):
                    raise PlanilhaInvalidaError("identity")
                normalized[field] = value
            elif field == "casa":
                normalized[field] = casa_canonica(_texto(value))
                if normalized[field] is None:
                    raise PlanilhaInvalidaError("house")
            elif field in {"odd", "stake_unidades"}:
                if isinstance(value, bool):
                    raise PlanilhaInvalidaError("boolean_number")
                decimal = value
                if isinstance(value, str):
                    decimal = value.strip().replace("R$", "").strip()
                    if "," in decimal:
                        decimal = decimal.replace(".", "").replace(",", ".")
                number = _numero(decimal, row["number"], field)
                if (field == "odd" and not 1.01 <= number <= 1000) or (
                    field == "stake_unidades" and number < 0
                ):
                    raise PlanilhaInvalidaError("range")
                normalized[field] = number
            elif field in {"data_aposta", "data_jogo", "atualizada_em"}:
                normalized[field] = (
                    None
                    if field == "data_jogo" and value in {None, ""}
                    else _data(value, row["number"], field).isoformat()
                )
            elif field == "freebet":
                normalized[field] = _booleano(value, row["number"])
            elif field == "conta_casa_ref":
                normalized[field] = _account(value, row["number"])
            else:
                if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
                    raise PlanilhaInvalidaError("unsafe_text")
                normalized[field] = _texto(value)
        except (ValueError, TypeError, OverflowError):
            errors.append({"field": field, "code": "invalid_value"})
    if row["formulas"] and not any(e["code"] == "formula_or_cell_error" for e in errors):
        errors.append({"field": None, "code": "formula_or_cell_error"})
    return normalized, errors
