"""Parse the documented Excel import template without executing workbook formulas."""

from __future__ import annotations

import hashlib
import math
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO
from typing import cast
from zipfile import BadZipFile, ZipFile
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from bancaemdia.core.body_limits import PLANILHA_MAX_BYTES
from bancaemdia.domain.materializar import casa_canonica

MAX_BYTES = PLANILHA_MAX_BYTES
MAX_ROWS = 1_000
REQUIRED = ("casa", "data_aposta", "odd", "stake_unidades", "atualizada_em")
OPTIONAL = ("evento", "descricao", "mercado_bruto", "freebet")
FUSO_DO_BRASIL = ZoneInfo("America/Sao_Paulo")


class PlanilhaInvalidaError(ValueError):
    pass


@dataclass(frozen=True)
class LinhaPlanilha:
    chave: str
    numero: int
    casa: str
    data_aposta: datetime
    atualizada_em: datetime
    odd: float
    stake_unidades: float
    evento: str | None
    descricao: str | None
    mercado_bruto: str | None
    freebet: bool
    presentes: frozenset[str]


def _data(valor: object, numero: int, campo: str) -> datetime:
    if isinstance(valor, datetime):
        data = valor
    elif isinstance(valor, date):
        data = datetime.combine(valor, time.min)
    elif isinstance(valor, str):
        try:
            data = datetime.fromisoformat(valor.strip())
        except ValueError as erro:
            raise PlanilhaInvalidaError(f"linha {numero}: {campo} inválida") from erro
    else:
        raise PlanilhaInvalidaError(f"linha {numero}: {campo} ausente")
    return data if data.tzinfo else data.replace(tzinfo=FUSO_DO_BRASIL)


def _numero(valor: object, numero: int, campo: str) -> float:
    try:
        convertido = float(cast(str | float | int, valor))
    except (TypeError, ValueError) as erro:
        raise PlanilhaInvalidaError(f"linha {numero}: {campo} inválida") from erro
    if not math.isfinite(convertido):
        raise PlanilhaInvalidaError(f"linha {numero}: {campo} inválida")
    return convertido


def _texto(valor: object) -> str | None:
    texto = str(valor).strip() if valor is not None else ""
    return texto or None


def _booleano(valor: object, numero: int) -> bool:
    if valor in (None, "", False, 0):
        return False
    if valor in (True, 1):
        return True
    if isinstance(valor, str):
        if valor.strip().lower() in {"sim", "true", "1"}:
            return True
        if valor.strip().lower() in {"não", "nao", "false", "0"}:
            return False
    raise PlanilhaInvalidaError(f"linha {numero}: freebet inválida")


def _campo(colunas: dict[str, int], valores: tuple[object, ...], nome: str) -> object:
    indice = colunas.get(nome)
    return None if indice is None or indice >= len(valores) else valores[indice]


def ler_planilha(conteudo: bytes, origem_id: str) -> list[LinhaPlanilha]:
    if not conteudo or len(conteudo) > MAX_BYTES:
        raise PlanilhaInvalidaError("arquivo Excel vazio ou maior que 5 MB")
    if not origem_id.strip() or len(origem_id) > 128:
        raise PlanilhaInvalidaError(
            "origem_id deve identificar a mesma planilha em todas as importações"
        )
    try:
        with ZipFile(BytesIO(conteudo)) as pacote:
            if sum(item.file_size for item in pacote.infolist()) > 50_000_000:
                raise PlanilhaInvalidaError("arquivo Excel descompactado maior que 50 MB")
    except BadZipFile as erro:
        raise PlanilhaInvalidaError("arquivo Excel inválido") from erro
    try:
        livro = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as erro:
        raise PlanilhaInvalidaError("arquivo Excel inválido") from erro
    linhas: list[LinhaPlanilha] = []
    try:
        for aba in livro.worksheets:
            iterador = aba.iter_rows(values_only=True)
            cabecalho = next(iterador, None)
            if cabecalho is None:
                continue
            colunas = {
                str(valor).strip().lower(): indice
                for indice, valor in enumerate(cabecalho)
                if valor is not None
            }
            if any(campo not in colunas for campo in REQUIRED):
                raise PlanilhaInvalidaError(f"aba {aba.title}: faltam colunas obrigatórias")
            for numero, valores in enumerate(iterador, start=2):
                if not any(valor is not None for valor in valores):
                    continue
                if len(linhas) >= MAX_ROWS:
                    raise PlanilhaInvalidaError("a importação aceita até 1.000 linhas")

                casa = casa_canonica(_texto(_campo(colunas, valores, "casa")))
                if casa is None:
                    raise PlanilhaInvalidaError(f"linha {numero}: casa não reconhecida")
                odd = _numero(_campo(colunas, valores, "odd"), numero, "odd")
                stake = _numero(
                    _campo(colunas, valores, "stake_unidades"), numero, "stake_unidades"
                )
                if odd < 1.01 or odd > 1000 or stake < 0:
                    raise PlanilhaInvalidaError(f"linha {numero}: odd ou stake fora do intervalo")
                identidade = f"{origem_id.strip()}:{aba.title}:{numero}".encode()
                linhas.append(
                    LinhaPlanilha(
                        chave=f"s:{hashlib.sha256(identidade).hexdigest()}",
                        numero=numero,
                        casa=casa,
                        data_aposta=_data(
                            _campo(colunas, valores, "data_aposta"), numero, "data_aposta"
                        ),
                        atualizada_em=_data(
                            _campo(colunas, valores, "atualizada_em"), numero, "atualizada_em"
                        ),
                        odd=odd,
                        stake_unidades=stake,
                        evento=_texto(_campo(colunas, valores, "evento")),
                        descricao=_texto(_campo(colunas, valores, "descricao")),
                        mercado_bruto=_texto(_campo(colunas, valores, "mercado_bruto")),
                        freebet=_booleano(_campo(colunas, valores, "freebet"), numero),
                        presentes=frozenset(colunas),
                    )
                )
    finally:
        livro.close()
    return linhas
