"""Authenticated access to the account's exportable records."""

from io import BytesIO
from typing import Annotated, Any, Literal

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.encoders import jsonable_encoder
from fastapi.responses import JSONResponse, Response
from openpyxl import Workbook
from pydantic import BaseModel, ConfigDict
from sqlalchemy import delete, select, update
from sqlalchemy.ext.asyncio import AsyncSession

from bancaemdia.api.contracts import AUTHENTICATED_ERROR_RESPONSES, ErrorResponse
from bancaemdia.api.deps import get_current_user, get_current_user_snapshot
from bancaemdia.db.session import get_db, get_db_snapshot
from bancaemdia.domain.registros import Usuario
from bancaemdia.models import Aposta, Base, Evento, RevisaoPendente, Upload
from bancaemdia.models.usuario import Usuario as UsuarioModel

router = APIRouter(responses=AUTHENTICATED_ERROR_RESPONSES)

# Binary uploads and credential hashes are deliberately excluded. Shared Telegram
# messages/photos have no tenant owner in the present schema and cannot be safely
# attributed to one account; see docs/SECURITY.md before calling this a full export.
EXPORT_TABLES = (
    "apostas",
    "eventos",
    "movimentos",
    "movimento_requisicoes",
    "bancas",
    "contas_casa",
    "unidades",
    "revisao_pendente",
    "coletas_casa",
    "uploads",
    "upload_bilhetes",
    "upload_arquivos",
    "coleta_token",
    "chamadas_ia",
    "audit_log",
)
EXCLUDED_COLUMNS = frozenset({"token_hash", "conteudo"})
EXPORT_RESPONSE_SCHEMA: dict[str, Any] = {
    "type": "object",
    "required": ["usuario", *EXPORT_TABLES],
    "properties": {
        "usuario": {
            "type": "object",
            "additionalProperties": {"$ref": "#/components/schemas/JsonValue"},
        },
        **{
            name: {
                "type": "array",
                "items": {
                    "type": "object",
                    "additionalProperties": {"$ref": "#/components/schemas/JsonValue"},
                },
            }
            for name in EXPORT_TABLES
        },
    },
    "additionalProperties": False,
}


class AnonymizeResponse(BaseModel):
    model_config = ConfigDict(extra="forbid")

    status: Literal["anonymized"]


async def collect_user_data(session: AsyncSession, usuario: Usuario) -> dict[str, Any]:
    data: dict[str, Any] = {
        "usuario": jsonable_encoder({
            "id": usuario.id,
            "email": usuario.email,
            "nome": usuario.nome,
            "criado_em": usuario.criado_em,
            "ativo": usuario.ativo,
        })
    }
    for name in EXPORT_TABLES:
        table = Base.metadata.tables[name]
        columns = [column for column in table.columns if column.name not in EXCLUDED_COLUMNS]
        statement = select(*columns).where(table.c.usuario_id == usuario.id)
        statement = statement.order_by(*(column for column in table.primary_key.columns))
        rows = await session.execute(statement)
        data[name] = [jsonable_encoder(dict(row._mapping)) for row in rows]
    return data


def _safe_excel_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        import json

        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def _excel(data: dict[str, Any]) -> bytes:
    workbook = Workbook(write_only=True)
    account = workbook.create_sheet("usuario")
    account.append(list(data["usuario"]))
    account.append([_safe_excel_value(value) for value in data["usuario"].values()])
    for name in EXPORT_TABLES:
        sheet = workbook.create_sheet(name)
        table = Base.metadata.tables[name]
        columns = [column.name for column in table.columns if column.name not in EXCLUDED_COLUMNS]
        sheet.append(columns)
        for row in data[name]:
            sheet.append([_safe_excel_value(row[column]) for column in columns])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@router.get(
    "/api/v1/usuario/me/export",
    summary="Exportar meus registros",
    description="Baixa os registros vinculados à conta em JSON ou Excel. Não inclui arquivos binários.",
    responses={
        200: {
            "description": "Exportação dos registros vinculados à conta.",
            "content": {
                "application/json": {"schema": EXPORT_RESPONSE_SCHEMA},
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {
                    "schema": {"type": "string", "format": "binary"}
                },
            },
        }
    },
)
async def exportar_meus_dados(
    usuario: Annotated[Usuario, Depends(get_current_user_snapshot)],
    session: Annotated[AsyncSession, Depends(get_db_snapshot)],
    formato: Annotated[Literal["json", "xlsx"], Query(description="Formato do arquivo")] = "json",
) -> Response:
    data = await collect_user_data(session, usuario)
    headers = {
        "Cache-Control": "no-store",
        "Content-Disposition": f'attachment; filename="meus-dados.{formato}"',
    }
    if formato == "xlsx":
        return Response(
            _excel(data),
            headers=headers,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    return JSONResponse(data, headers=headers)


@router.delete(
    "/api/v1/usuario/me",
    summary="Desativar e anonimizar minha conta",
    description=(
        "Desativa a conta, remove registros vinculados e esvazia o conteúdo dos eventos "
        "mantidos para auditoria. Contas com mensagens ou fotos importadas exigem "
        "tratamento separado e recebem 409, sem qualquer alteração."
    ),
    response_model=AnonymizeResponse,
    responses={
        409: {
            "model": ErrorResponse,
            "description": "Há mensagens ou fotos sem dono exclusivo no banco.",
        }
    },
)
async def anonimizar_minha_conta(
    usuario: Annotated[Usuario, Depends(get_current_user)],
    session: Annotated[AsyncSession, Depends(get_db)],
) -> JSONResponse:
    # Serializa pedidos concorrentes de exclusão. O trigger de usuário inativo
    # impede novas gravações por trabalhadores que ainda tenham tarefas antigas.
    await session.execute(
        select(UsuarioModel.id).where(UsuarioModel.id == usuario.id).with_for_update()
    )
    raw_sources = (
        select(Upload.id).where(Upload.usuario_id == usuario.id, Upload.chat_id.is_not(None)),
        select(Aposta.id).where(
            Aposta.usuario_id == usuario.id,
            (Aposta.chat_id.is_not(None) | Aposta.midia_hash.is_not(None)),
        ),
        select(RevisaoPendente.id).where(
            RevisaoPendente.usuario_id == usuario.id,
            RevisaoPendente.midia_hash.is_not(None),
        ),
    )
    for statement in raw_sources:
        if await session.scalar(statement.limit(1)) is not None:
            raise HTTPException(
                status_code=409,
                detail="Há mensagens ou fotos importadas que exigem exclusão assistida.",
            )
    for table in reversed(Base.metadata.sorted_tables):
        if table.name in {"usuarios", "eventos", "audit_log"} or "usuario_id" not in table.c:
            continue
        await session.execute(delete(table).where(table.c.usuario_id == usuario.id))
    await session.execute(
        update(Evento)
        .where(Evento.usuario_id == usuario.id)
        .values(payload_json={}, chat_id=None, message_id=None, aposta_chave=None, confianca=None)
    )
    await session.execute(
        update(UsuarioModel)
        .where(UsuarioModel.id == usuario.id)
        .values(
            email=f"deleted-{usuario.id}@invalid.local",
            nome="Conta excluída",
            ativo=False,
        )
    )
    await session.commit()
    return JSONResponse(
        {"status": "anonymized"},
        headers={"Cache-Control": "no-store"},
    )
